from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document as WordDocument
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend import api
from backend.db import Base, Document, ObjectGroup, ObjectRecord, Template
from backend.schemas import DocumentPatch


class BackendApiRegressionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.session_factory = sessionmaker(bind=self.engine, autoflush=False, expire_on_commit=False)
        self.session_patch = patch.object(api, "SessionLocal", self.session_factory)
        self.settings_patch = patch.object(
            api,
            "settings",
            SimpleNamespace(vault_dir=self.root / "vault", frontend_dir=self.root / "frontend"),
        )
        self.session_patch.start()
        self.settings_patch.start()

    def tearDown(self) -> None:
        self.settings_patch.stop()
        self.session_patch.stop()
        self.engine.dispose()
        self.temporary_directory.cleanup()

    def create_object(self, *, customer: str = "CURRENT CUSTOMER") -> int:
        with self.session_factory.begin() as db:
            obj = ObjectRecord(
                code="TEST-001",
                name="Test object",
                short_name="Test",
                address="Samara",
                folder_path="objects/TEST-001",
                is_current=True,
                profile="industrial_repair",
                meta_json={"customer": {"name": customer}},
            )
            db.add(obj)
            db.flush()
            return obj.id

    def create_group(self, object_id: int, *, title: str = "Acts") -> int:
        with self.session_factory.begin() as db:
            group = ObjectGroup(
                object_id=object_id,
                code=f"group-{title.lower()}",
                title=title,
                kind="acts",
                sort_order=10,
                is_required=True,
            )
            db.add(group)
            db.flush()
            return group.id

    def create_template(self, fields: list[dict], *, marker_map: dict | None = None, template_file: str = "") -> int:
        with self.session_factory.begin() as db:
            template = Template(
                title="Test template",
                doc_type="aosr",
                group_kind="acts",
                fields_json=fields,
                field_map_json={},
                marker_map=marker_map or {},
                template_file=template_file,
            )
            db.add(template)
            db.flush()
            return template.id

    def create_document(
        self,
        object_id: int,
        *,
        template_id: int | None = None,
        group_id: int | None = None,
        completion_status: str = "not_started",
        values: dict | None = None,
        missing: list | None = None,
    ) -> int:
        with self.session_factory.begin() as db:
            document = Document(
                object_id=object_id,
                template_id=template_id,
                group_id=group_id,
                number="A-1",
                title="Test document",
                doc_type="aosr",
                work_title="Test work",
                document_date="27.07.2026",
                status="draft",
                completion_status=completion_status,
                values_json=values or {},
                warnings_json=[],
                missing_fields_json=missing or [],
            )
            db.add(document)
            db.flush()
            return document.id

    def test_object_fields_are_fresh_on_read_and_generation(self) -> None:
        object_id = self.create_object(customer="NEW CUSTOMER")
        source = self.root / "source.docx"
        word = WordDocument()
        word.add_paragraph("Customer: [[customer]]; Note: [[note]]")
        word.save(source)
        template_id = self.create_template(
            [
                {"key": "customer.name", "source": "object.meta.customer.name", "required": True},
                {"key": "manual.note", "source": "manual", "required": True},
            ],
            marker_map={
                "customer": {"field": "customer.name"},
                "note": {"field": "manual.note"},
            },
            template_file="source.docx",
        )
        document_id = self.create_document(
            object_id,
            template_id=template_id,
            values={"customer.name": "OLD CUSTOMER", "manual.note": "KEEP MANUAL"},
        )

        read_documents = [
            api.object_workspace(object_id)["documents"][0],
            api.app_state()["documents"][0],
            api.list_documents(object_id=object_id)[0],
            api.global_search("Test")["documents"][0],
        ]
        for document in read_documents:
            values = document["values_json"]
            self.assertEqual("NEW CUSTOMER", values["customer.name"])
            self.assertEqual("KEEP MANUAL", values["manual.note"])
        with self.session_factory() as db:
            self.assertEqual("OLD CUSTOMER", db.get(Document, document_id).values_json["customer.name"])

        with patch.object(api, "resolve_template_path", return_value=source):
            result = api.generate_documents({"ids": [document_id]})

        output = Path(result["made"][0]["file"])
        rendered = WordDocument(output)
        rendered_text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
        self.assertIn("NEW CUSTOMER", rendered_text)
        self.assertNotIn("OLD CUSTOMER", rendered_text)
        self.assertIn("KEEP MANUAL", rendered_text)
        with self.session_factory() as db:
            persisted = db.get(Document, document_id)
            self.assertEqual("NEW CUSTOMER", persisted.values_json["customer.name"])
            self.assertEqual("KEEP MANUAL", persisted.values_json["manual.note"])

    def test_empty_object_field_hides_obsolete_snapshot_on_read(self) -> None:
        object_id = self.create_object(customer="")
        template_id = self.create_template(
            [{"key": "customer.name", "source": "object.meta.customer.name", "required": True}]
        )
        self.create_document(
            object_id,
            template_id=template_id,
            values={"customer.name": "OBSOLETE CUSTOMER"},
        )

        document = api.object_workspace(object_id)["documents"][0]

        self.assertNotIn("customer.name", document["values_json"])
        self.assertEqual(["customer.name"], document["missing_fields_json"])

    def test_content_patch_honors_explicit_completion_status(self) -> None:
        object_id = self.create_object()
        template_id = self.create_template(
            [{"key": "manual.value", "source": "manual", "required": True}]
        )
        document_id = self.create_document(
            object_id,
            template_id=template_id,
            values={},
            missing=["manual.value"],
        )

        updated = api.update_document(
            document_id,
            DocumentPatch(values_json={"manual.value": "Filled"}, completion_status="completed"),
        )

        self.assertEqual("completed", updated.completion_status)
        self.assertEqual([], updated.missing_fields_json)

        revised = api.update_document(
            document_id,
            DocumentPatch(values_json={"manual.value": "Changed"}),
        )
        self.assertEqual("needs_revision", revised.completion_status)

    def test_completeness_sections_require_completed_and_no_missing_fields(self) -> None:
        object_id = self.create_object()
        group_id = self.create_group(object_id)
        self.create_document(object_id, group_id=group_id, completion_status="completed")
        self.create_document(object_id, group_id=group_id, completion_status="not_started")
        self.create_document(
            object_id,
            group_id=group_id,
            completion_status="completed",
            missing=["legacy.required"],
        )

        report = api.completeness_report()

        self.assertEqual(1, report["ready"])
        self.assertEqual(1, report["sections"][0]["ready"])
        self.assertEqual(3, report["sections"][0]["total"])
        self.assertEqual("warning", report["sections"][0]["status"])


if __name__ == "__main__":
    unittest.main()
