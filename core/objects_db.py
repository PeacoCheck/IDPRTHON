# -*- coding: utf-8 -*-
"""Excel registry of local construction objects."""
from __future__ import annotations

import datetime as dt
import re
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .excel_templates import copy_template_or_create_new
from .paths import (
    AOSR_FILE_NAME,
    LEGACY_AOSR_DATA_FILE,
    OBJECTS_DB,
    OBJECTS_SHEET,
    ensure_base_dirs,
    object_folder,
    rel_to_base,
    resolve_project_path,
    resolve_data_file_for_object,
)

HEADERS = [
    "id",
    "Наименование объекта",
    "Адрес объекта",
    "Файл данных",
    "Папка объекта",
    "Статус",
    "Текущий",
    "Дата создания",
    "Комментарий",
]


def ensure_objects_db() -> None:
    ensure_base_dirs()
    if not OBJECTS_DB.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = OBJECTS_SHEET
        ws.append(HEADERS)
        if LEGACY_AOSR_DATA_FILE.exists():
            _append_legacy_object(ws)
        wb.save(OBJECTS_DB)
        return

    wb = load_workbook(OBJECTS_DB)
    ws = wb[OBJECTS_SHEET] if OBJECTS_SHEET in wb.sheetnames else wb.create_sheet(OBJECTS_SHEET, 0)
    for col, header in enumerate(HEADERS, start=1):
        if ws.cell(row=1, column=col).value != header:
            ws.cell(row=1, column=col, value=header)
    wb.save(OBJECTS_DB)


def load_objects() -> list[dict]:
    ensure_objects_db()
    wb = load_workbook(OBJECTS_DB, data_only=True)
    ws = wb[OBJECTS_SHEET]
    objects = []
    for row in ws.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
        object_id = _text(row[0])
        if not object_id:
            continue
        name = _text(row[1])
        obj = {
            "id": object_id,
            "name": name,
            "shortName": _short_name(name or object_id),
            "address": _text(row[2]),
            "dataFile": _text(row[3]),
            "folder": _text(row[4]),
            "status": _text(row[5]) or "active",
            "current": _text(row[6]).lower() == "yes",
            "created": _text(row[7]),
            "comment": _text(row[8]),
        }
        objects.append(obj)
    return objects


def get_current_object() -> dict | None:
    objects = load_objects()
    for obj in objects:
        if obj["status"] == "active" and obj["current"]:
            return obj
    for obj in objects:
        if obj["status"] == "active":
            return obj
    return None


def set_current_object(object_id: str) -> None:
    wb, ws = _open_db()
    found = False
    for row in range(2, ws.max_row + 1):
        row_id = _text(ws.cell(row=row, column=1).value)
        status = _text(ws.cell(row=row, column=6).value) or "active"
        ws.cell(row=row, column=7, value="")
        if row_id == object_id and status == "active":
            ws.cell(row=row, column=7, value="yes")
            found = True
    if not found:
        raise ValueError("Активный объект не найден.")
    wb.save(OBJECTS_DB)


def add_object(data: dict) -> dict:
    name = _text(data.get("name"))
    if not name:
        raise ValueError("Введите наименование объекта.")
    address = _text(data.get("address"))
    comment = _text(data.get("comment"))

    wb, ws = _open_db()
    object_id = _next_object_id(ws)
    folder = object_folder(object_id)
    data_file = folder / AOSR_FILE_NAME
    copy_template_or_create_new(data_file, name, address)

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=7, value="")
    ws.append([
        object_id,
        name,
        address,
        rel_to_base(data_file),
        rel_to_base(folder),
        "active",
        "yes",
        dt.datetime.now().strftime("%d.%m.%Y"),
        comment,
    ])
    wb.save(OBJECTS_DB)
    return get_current_object()


def archive_object(object_id: str) -> None:
    wb, ws = _open_db()
    found = False
    was_current = False
    for row in range(2, ws.max_row + 1):
        if _text(ws.cell(row=row, column=1).value) == object_id:
            found = True
            was_current = _text(ws.cell(row=row, column=7).value).lower() == "yes"
            ws.cell(row=row, column=6, value="archive")
            ws.cell(row=row, column=7, value="")
            break
    if not found:
        raise ValueError("Объект не найден.")
    if was_current:
        _mark_first_active_current(ws)
    wb.save(OBJECTS_DB)


def unarchive_object(object_id: str) -> None:
    wb, ws = _open_db()
    found = False
    for row in range(2, ws.max_row + 1):
        if _text(ws.cell(row=row, column=1).value) == object_id:
            ws.cell(row=row, column=6, value="active")
            found = True
            break
    if not found:
        raise ValueError("Объект не найден.")
    if not _has_current_active(ws):
        _mark_first_active_current(ws)
    wb.save(OBJECTS_DB)


def get_object_by_id(object_id: str) -> dict | None:
    return next((obj for obj in load_objects() if obj["id"] == object_id), None)


def ensure_object_files(obj: dict) -> None:
    raw_folder = Path(obj.get("folder") or "")
    folder = raw_folder.resolve() if raw_folder.is_absolute() else resolve_project_path(raw_folder or f"data/objects/{obj['id']}")
    folder.mkdir(parents=True, exist_ok=True)
    data_file = resolve_data_file_for_object(obj)
    if not data_file.exists():
        copy_template_or_create_new(data_file, obj.get("name", ""), obj.get("address", ""))


def _append_legacy_object(ws) -> None:
    name, address = _read_project_name_address(LEGACY_AOSR_DATA_FILE)
    folder = object_folder("object_001")
    data_file = folder / AOSR_FILE_NAME
    folder.mkdir(parents=True, exist_ok=True)
    if not data_file.exists():
        shutil.copy2(LEGACY_AOSR_DATA_FILE, data_file)
    ws.append([
        "object_001",
        name or "Объект",
        address,
        rel_to_base(data_file),
        rel_to_base(folder),
        "active",
        "yes",
        dt.datetime.now().strftime("%d.%m.%Y"),
        "основной объект",
    ])


def _read_project_name_address(path: Path) -> tuple[str, str]:
    try:
        wb = load_workbook(path, data_only=True)
        ws = wb["1.Объект"]
        data = {}
        for key, val in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            key = _text(key)
            if key:
                data[key] = _text(val)
        return data.get("Наименование объекта", ""), data.get("Адрес объекта", "")
    except Exception:
        return "", ""


def _open_db():
    ensure_objects_db()
    wb = load_workbook(OBJECTS_DB)
    return wb, wb[OBJECTS_SHEET]


def _next_object_id(ws) -> str:
    max_num = 0
    for row in range(2, ws.max_row + 1):
        match = re.match(r"object_(\d+)$", _text(ws.cell(row=row, column=1).value))
        if match:
            max_num = max(max_num, int(match.group(1)))
    return f"object_{max_num + 1:03d}"


def _mark_first_active_current(ws) -> None:
    marked = False
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=7, value="")
        if not marked and (_text(ws.cell(row=row, column=6).value) or "active") == "active":
            ws.cell(row=row, column=7, value="yes")
            marked = True


def _has_current_active(ws) -> bool:
    for row in range(2, ws.max_row + 1):
        status = _text(ws.cell(row=row, column=6).value) or "active"
        current = _text(ws.cell(row=row, column=7).value).lower() == "yes"
        if status == "active" and current:
            return True
    return False


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _short_name(name: str) -> str:
    return name if len(name) <= 32 else name[:32].rstrip() + "…"
