# -*- coding: utf-8 -*-
"""Creation and copying of Excel workbooks for object data."""
from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .paths import AOSR_TEMPLATE, LEGACY_AOSR_DATA_FILE

PROJECT_ROWS = [
    ("Наименование объекта", ""),
    ("Адрес объекта", ""),
    ("Застройщик (технический заказчик)", ""),
    ("Лицо, осуществляющее строительство", ""),
    ("Лицо, осуществляющее подготовку проектной документации", ""),
    ("Лицо, выполнившее работы (по умолчанию)", ""),
    ("Папка для готовых актов", "Акты/АОСР"),
    ("Количество экземпляров (по умолчанию)", "2"),
    ("Шаблон имени файла", "АОСР_{номер}_{дата}"),
]


def create_empty_aosr_workbook(path: str | Path, object_name: str, object_address: str) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "1.Объект"
    ws.append(["Параметр", "Значение"])
    for key, default in PROJECT_ROWS:
        value = object_name if key == "Наименование объекта" else object_address if key == "Адрес объекта" else default
        ws.append([key, value])

    people = wb.create_sheet("2.Участники")
    people.append(["Роль", "Организация", "Должность", "ФИО", "Документ", "НРС"])

    materials = wb.create_sheet("3.Материалы")
    materials.append(["Код", "Наименование", "Тип документа", "Номер", "Дата", "Примечание"])

    acts = wb.create_sheet("4.Реестр актов")
    acts.append([
        "№ акта", "Дата акта", "Наименование работ", "Проектная документация",
        "Материалы", "Предъявленные документы", "Начало работ", "Окончание работ",
        "Нормативные документы", "Последующие работы", "Дополнительно",
        "Экземпляров", "Приложения", "Исполнитель",
    ])

    wb.save(path)


def copy_template_or_create_new(path: str | Path, object_name: str, object_address: str) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    template = AOSR_TEMPLATE if AOSR_TEMPLATE.exists() else LEGACY_AOSR_DATA_FILE
    if template.exists():
        shutil.copy2(template, path)
        update_project_sheet(path, object_name, object_address)
    else:
        create_empty_aosr_workbook(path, object_name, object_address)


def update_project_sheet(path: str | Path, object_name: str, object_address: str) -> None:
    wb = load_workbook(path)
    if "1.Объект" not in wb.sheetnames:
        wb.create_sheet("1.Объект", 0)
    ws = wb["1.Объект"]
    if ws.max_row < 1:
        ws.append(["Параметр", "Значение"])

    values = {
        "Наименование объекта": object_name,
        "Адрес объекта": object_address,
        "Папка для готовых актов": "Акты/АОСР",
        "Количество экземпляров (по умолчанию)": "2",
        "Шаблон имени файла": "АОСР_{номер}_{дата}",
    }
    existing = {}
    for row in range(2, ws.max_row + 1):
        key = str(ws.cell(row=row, column=1).value or "").strip()
        if key:
            existing[key] = row

    for key, value in values.items():
        row = existing.get(key)
        if row is None:
            ws.append([key, value])
        else:
            ws.cell(row=row, column=2, value=value)
    wb.save(path)
