from __future__ import annotations

import os
import json
import mimetypes
import re
import shutil
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from docx import Document as WordDocument
from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import delete, func, or_, select

from .config import BASE_DIR, settings
from .db import (
    CatalogItem,
    CorpusSource,
    Document,
    ObjectGroup,
    ObjectRecord,
    Preset,
    Section,
    SessionLocal,
    StructureException,
    StructureRule,
    Template,
)
from .docx_templates import inspect_docx_markers, render_docx_template, resolve_marker_values
from .schemas import (
    DocumentCreate,
    DocumentOut,
    DocumentPatch,
    ObjectCreate,
    ObjectOut,
    ObjectPatch,
    TemplateBind,
    TemplateCopy,
    TemplateDuplicate,
    TemplatePatch,
)
from .seed import GROUPS

router = APIRouter(prefix="/api")
MAX_TEMPLATE_SIZE = 25 * 1024 * 1024
MAX_DOCUMENT_SIZE = 100 * 1024 * 1024
MAX_SPREADSHEET_PREVIEW_SIZE = 40 * 1024 * 1024
MAX_SPREADSHEET_EXPANDED_SIZE = 200 * 1024 * 1024
MAX_SPREADSHEET_ARCHIVE_ENTRIES = 5000
TEMPLATE_SUFFIXES = {".docx", ".xlsx", ".pdf"}
DOCUMENT_SUFFIXES = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".xlsm", ".jpg", ".jpeg", ".png", ".webp", ".dwg", ".dxf", ".txt"}
MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xls": "application/vnd.ms-excel",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
    ".dwg": "application/acad",
    ".dxf": "application/dxf",
    ".txt": "text/plain",
}
DOCUMENT_GROUP_KIND_BY_TYPE = {
    "registry": "registry",
    "general_journal": "journals",
    "incoming_journal": "journals",
    "welding_journal": "journals",
    "appointment_order": "permits",
    "aosr": "acts",
    "vik": "acts",
    "hydro": "acts",
    "dismantling_act": "acts",
    "installation_act": "acts",
    "equipment_acceptance": "acts",
    "repair_readiness": "acts",
    "certificate": "quality",
    "passport": "quality",
    "incoming_control": "quality",
    "executive_scheme": "schemes",
    "project": "working",
    "working_documentation": "working",
    "ppr": "source",
    "method_statement": "source",
}
STANDARD_GROUP_CODES = {code for code, *_rest in GROUPS}


def safe_component(value: str, fallback: str = "Документ") -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value)).strip(" .")
    return (cleaned[:120] or fallback).strip()


def safe_file_name(value: str, fallback: str = "Документ") -> str:
    raw_name = Path(str(value or fallback)).name
    suffix = Path(raw_name).suffix
    stem = safe_component(Path(raw_name).stem, fallback)
    max_stem_length = max(1, 120 - len(suffix))
    safe_stem = stem[:max_stem_length].rstrip(" .") or fallback
    return f"{safe_stem}{suffix}"


def is_within(path: Path, root: Path) -> bool:
    path = path.resolve()
    root = root.resolve()
    return path == root or root in path.parents


def media_type_for_path(path: Path) -> str:
    suffix = path.suffix.casefold()
    return MEDIA_TYPES.get(suffix) or mimetypes.guess_type(path.name)[0] or "application/octet-stream"


def safe_vault_file(path_value: str | Path | None, *, must_exist: bool = True) -> Path | None:
    if not path_value:
        return None
    candidate = Path(path_value).expanduser().resolve()
    vault_root = settings.vault_dir.resolve()
    if not is_within(candidate, vault_root):
        return None
    if must_exist and not candidate.is_file():
        return None
    return candidate


def file_info_for_path(
    path_value: str | Path | None,
    *,
    download_url: str = "",
    preview_url: str = "",
    allowed_roots: tuple[Path, ...] | None = None,
) -> dict[str, Any] | None:
    if allowed_roots is None:
        path = safe_vault_file(path_value)
    elif path_value:
        candidate = Path(path_value).expanduser().resolve()
        path = candidate if candidate.is_file() and any(is_within(candidate, root) for root in allowed_roots) else None
    else:
        path = None
    if not path:
        return None
    stat = path.stat()
    return {
        "file_name": path.name,
        "file_size": stat.st_size,
        "file_extension": path.suffix.casefold(),
        "media_type": media_type_for_path(path),
        "updated_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
        "download_url": download_url,
        "preview_url": preview_url,
    }


def write_upload_limited(upload: UploadFile, target: Path, limit: int, message: str) -> int:
    target.parent.mkdir(parents=True, exist_ok=True)
    written = 0
    try:
        with target.open("wb") as output:
            while True:
                chunk = upload.file.read(1024 * 1024)
                if not chunk:
                    break
                written += len(chunk)
                if written > limit:
                    raise HTTPException(413, message)
                output.write(chunk)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    return written


def template_roots() -> tuple[Path, Path]:
    return ((BASE_DIR / "resources" / "templates").resolve(), (settings.vault_dir / "templates").resolve())


def resolve_template_path(value: str, *, must_exist: bool = True) -> Path:
    raw = Path(value).expanduser()
    candidate = raw.resolve() if raw.is_absolute() else (BASE_DIR / raw).resolve()
    if candidate.suffix.casefold() not in TEMPLATE_SUFFIXES or not any(is_within(candidate, root) for root in template_roots()):
        raise HTTPException(422, "Шаблон DOCX, XLSX или PDF должен находиться в resources/templates или vault/templates")
    if must_exist and not candidate.is_file():
        raise HTTPException(404, "Файл шаблона не найден")
    return candidate


def stored_template_path(path: Path) -> str:
    resolved = path.resolve()
    if is_within(resolved, BASE_DIR):
        return resolved.relative_to(BASE_DIR.resolve()).as_posix()
    return str(resolved)


def managed_template_target(template: Template, original_name: str) -> tuple[Path, int]:
    root = (settings.vault_dir / "templates").resolve()
    root.mkdir(parents=True, exist_ok=True)
    version = template.version + (1 if template.template_file else 0)
    stem = safe_component(Path(original_name).stem, "template")
    suffix = Path(original_name).suffix.casefold()
    if suffix not in TEMPLATE_SUFFIXES:
        raise HTTPException(415, "Шаблон должен быть файлом DOCX, XLSX или PDF")
    target = root / f"template-{template.id}-v{version}-{stem}{suffix}"
    while target.exists() and stored_template_path(target) != template.template_file:
        version += 1
        target = root / f"template-{template.id}-v{version}-{stem}{suffix}"
    return target, version


def validate_template_file(path: Path) -> None:
    suffix = path.suffix.casefold()
    try:
        if suffix == ".docx":
            inspect_docx_markers(path)
        elif suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            workbook.close()
        elif suffix == ".pdf":
            reader = PdfReader(str(path))
            len(reader.pages)
        else:
            raise HTTPException(415, "Шаблон должен быть файлом DOCX, XLSX или PDF")
    except HTTPException:
        raise
    except Exception as exc:
        label = suffix.removeprefix(".").upper() or "файл"
        raise HTTPException(422, f"Файл не является читаемым {label}: {exc}") from exc


def spreadsheet_preview_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)[:5000]


def spreadsheet_preview_data(path: Path) -> dict[str, Any]:
    suffix = path.suffix.casefold()
    if suffix not in {".xlsx", ".xlsm"}:
        raise HTTPException(415, "Предпросмотр доступен для XLSX и XLSM")
    if path.stat().st_size > MAX_SPREADSHEET_PREVIEW_SIZE:
        raise HTTPException(413, "Таблица больше 40 МБ — скачайте её для просмотра")
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_SPREADSHEET_ARCHIVE_ENTRIES:
                raise HTTPException(422, "В таблице слишком много внутренних файлов")
            if sum(entry.file_size for entry in entries) > MAX_SPREADSHEET_EXPANDED_SIZE:
                raise HTTPException(413, "Распакованный размер таблицы слишком велик для предпросмотра")
    except HTTPException:
        raise
    except (BadZipFile, OSError) as exc:
        raise HTTPException(422, "Файл не является читаемой таблицей XLSX") from exc

    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        if not workbook.worksheets:
            return {"sheet_name": "Лист 1", "rows": [], "truncated_rows": False, "truncated_columns": False}
        sheet = workbook.worksheets[0]
        maximum_row = max(0, int(sheet.max_row or 0))
        maximum_column = max(0, int(sheet.max_column or 0))
        rows = [
            [spreadsheet_preview_value(value) for value in row]
            for row in sheet.iter_rows(
                min_row=1,
                max_row=min(maximum_row, 50),
                min_col=1,
                max_col=min(maximum_column, 30),
                values_only=True,
            )
        ] if maximum_row and maximum_column else []
        while rows and not any(value not in (None, "") for value in rows[-1]):
            rows.pop()
        return {
            "sheet_name": sheet.title,
            "rows": rows,
            "truncated_rows": maximum_row > 50,
            "truncated_columns": maximum_column > 30,
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Не удалось подготовить безопасный предпросмотр таблицы") from exc
    finally:
        if workbook is not None:
            workbook.close()


def template_markers_payload(template: Template) -> dict[str, Any]:
    if not template.template_file:
        return {
            "template_id": template.id,
            "template_file": "",
            "file_info": None,
            "marker_map": template.marker_map or {},
            "markers": [],
        }
    path = resolve_template_path(template.template_file)
    markers = inspect_docx_markers(path) if path.suffix.casefold() == ".docx" else []
    configured = template.marker_map or {}
    for marker in markers:
        marker["mapping"] = configured.get(marker["name"])
    return {
        "template_id": template.id,
        "template_file": path.name,
        "file_info": file_info_for_path(
            path,
            download_url=f"/api/templates/{template.id}/file",
            preview_url=f"/api/templates/{template.id}/preview",
            allowed_roots=template_roots(),
        ),
        "marker_map": configured,
        "markers": markers,
    }


def normalized_marker_map(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise HTTPException(422, "marker_map должен быть JSON-объектом")
    nested = payload.get("markers")
    return nested if isinstance(nested, dict) else payload


def document_marker_context(document: Document, obj: ObjectRecord) -> dict[str, Any]:
    values = dict(document.values_json or {})
    meta = dict(obj.meta_json or {})
    customer = dict(meta.get("customer") or {})
    contractor = dict(meta.get("contractor") or {})
    responsible = dict(meta.get("responsible_person") or {})
    context: dict[str, Any] = {
        **values,
        "document": {
            "id": document.id,
            "number": document.number,
            "date": document.document_date,
            "title": document.title,
            "work_title": document.work_title,
        },
        "object": {
            "id": obj.id,
            "code": obj.code,
            "name": obj.name,
            "address": obj.address,
            "meta": meta,
        },
        "customer": customer,
        "contractor": contractor,
        "responsible": responsible,
        "document.number": document.number,
        "document.date": document.document_date,
        "journal.number": document.number,
        "object.code": obj.code,
        "object.name": obj.name,
        "object.address": obj.address,
        "object.meta.work_kind": meta.get("work_kind") or values.get("work_kind"),
        "customer.name": values.get("customer.name") or customer.get("name"),
        "contractor.name": values.get("contractor.name") or contractor.get("name"),
        "responsible.fio": values.get("responsible.fio") or responsible.get("fio"),
        "journal_number": document.number,
        "journal_date": document.document_date,
        "construction_type": values.get("construction_type") or meta.get("work_kind") or values.get("work_kind"),
        "object_name": obj.name,
        "object_address": obj.address,
        "zakaz": values.get("customer.name") or customer.get("name"),
        "podryad": values.get("contractor.name") or contractor.get("name"),
    }
    return context


def expected_output_path(document_id: int, object_id: int) -> Path:
    return (settings.vault_dir / "output" / f"object-{object_id}" / f"document-{document_id}.docx").resolve()


def safe_owned_output(path_value: str, document_id: int, object_id: int) -> Path | None:
    if not path_value:
        return None
    candidate = Path(path_value).expanduser().resolve()
    output_root = (settings.vault_dir / "output").resolve()
    expected = expected_output_path(document_id, object_id)
    if candidate == expected and is_within(candidate, output_root):
        return candidate
    return None


def infer_document_type(filename: str) -> str:
    name = filename.lower()
    patterns = (
        ("реестр", "registry"),
        ("общий журнал", "general_journal"),
        ("входн", "incoming_control"),
        ("свар", "welding_journal"),
        ("аоср", "aosr"),
        ("скрыт", "aosr"),
        ("вик", "vik"),
        ("гидр", "hydro"),
        ("исполнительн", "executive_scheme"),
        ("сертификат", "certificate"),
        ("паспорт", "passport"),
        ("приказ", "appointment_order"),
        ("ппр", "ppr"),
    )
    return next((doc_type for marker, doc_type in patterns if marker in name), "imported")


def infer_group_kind(doc_type: str, title: str = "") -> str:
    normalized_type = str(doc_type or "").strip().casefold()
    if normalized_type in DOCUMENT_GROUP_KIND_BY_TYPE:
        return DOCUMENT_GROUP_KIND_BY_TYPE[normalized_type]

    inferred_type = infer_document_type(title)
    if inferred_type in DOCUMENT_GROUP_KIND_BY_TYPE:
        return DOCUMENT_GROUP_KIND_BY_TYPE[inferred_type]

    normalized_title = str(title or "").strip().casefold()
    title_patterns = (
        ("реестр", "registry"),
        ("журнал", "journals"),
        ("ппр", "source"),
        ("технолог", "source"),
        ("приказ", "permits"),
        ("допуск", "permits"),
        ("разреш", "permits"),
        ("схем", "schemes"),
        ("акт", "acts"),
        ("испыт", "acts"),
        ("сертификат", "quality"),
        ("паспорт", "quality"),
        ("материал", "quality"),
        ("проект", "working"),
        ("рабоч", "working"),
        ("фото", "evidence"),
        ("архив", "archive"),
    )
    return next((kind for marker, kind in title_patterns if marker in normalized_title), "other")


def auto_object_group(db, obj: ObjectRecord, doc_type: str, title: str = "") -> ObjectGroup | None:
    kind = infer_group_kind(doc_type, title)
    return db.scalar(
        select(ObjectGroup)
        .where(ObjectGroup.object_id == obj.id, ObjectGroup.kind == kind)
        .order_by(ObjectGroup.is_required.desc(), ObjectGroup.sort_order, ObjectGroup.id)
        .limit(1)
    )


def object_storage_root(obj: ObjectRecord) -> Path:
    root = (settings.vault_dir / obj.folder_path).resolve()
    vault_root = settings.vault_dir.resolve()
    if not is_within(root, vault_root):
        raise HTTPException(422, "Путь объекта выходит за пределы локального хранилища")
    return root


def unique_file_target(directory: Path, filename: str) -> Path:
    safe_name = safe_file_name(filename)
    target = directory / safe_name
    counter = 2
    while target.exists():
        target = directory / f"{Path(safe_name).stem}_{counter}{Path(safe_name).suffix}"
        counter += 1
    return target


def document_version_entries(document: Document) -> list[dict[str, Any]]:
    raw = (document.values_json or {}).get("versions", [])
    return [dict(item) for item in raw if isinstance(item, dict)] if isinstance(raw, list) else []


def archive_document_file(document: Document, obj: ObjectRecord) -> tuple[dict[str, Any] | None, Path | None, bool]:
    source = document_primary_file(document)
    if not source:
        return None, None, False
    object_root = object_storage_root(obj)
    versions_dir = object_root / "versions"
    versions_dir.mkdir(parents=True, exist_ok=True)
    versions = document_version_entries(document)
    known_versions = [int(item.get("version", 0)) for item in versions if str(item.get("version", "")).isdigit()]
    configured_version = (document.values_json or {}).get("file_version")
    version = int(configured_version) if str(configured_version or "").isdigit() else max(known_versions, default=0) + 1
    version = max(version, max(known_versions, default=0) + 1, 1)
    stem = safe_component(source.stem, "document")
    suffix = source.suffix.casefold()
    archived = versions_dir / f"{stem}-v{version}{suffix}"
    existing_ids = {str(item.get("id") or "") for item in versions}
    while archived.exists() or f"v{version}" in existing_ids:
        version += 1
        archived = versions_dir / f"{stem}-v{version}{suffix}"

    moved = is_within(source, object_root) and not is_within(source, versions_dir)
    if moved:
        shutil.move(str(source), str(archived))
    else:
        shutil.copy2(source, archived)
    stat = archived.stat()
    metadata = {
        "id": f"v{version}",
        "version": version,
        "name": source.name,
        "stored_name": archived.name,
        "size": stat.st_size,
        "file_extension": archived.suffix.casefold(),
        "media_type": media_type_for_path(archived),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "download_url": f"/api/documents/{document.id}/versions/v{version}",
    }
    return metadata, archived, moved


def resolve_document_version(document: Document, obj: ObjectRecord, version_id: str) -> tuple[dict[str, Any], Path]:
    normalized = version_id if version_id.startswith("v") else f"v{version_id}"
    metadata = next(
        (
            item
            for item in document_version_entries(document)
            if str(item.get("id") or "") == normalized or str(item.get("version") or "") == version_id.removeprefix("v")
        ),
        None,
    )
    if not metadata:
        raise HTTPException(404, "Версия документа не найдена")
    stored_name = str(metadata.get("stored_name") or "")
    if not stored_name or Path(stored_name).name != stored_name:
        raise HTTPException(404, "Файл версии документа не найден")
    versions_root = (object_storage_root(obj) / "versions").resolve()
    target = (versions_root / stored_name).resolve()
    if not is_within(target, versions_root) or not target.is_file():
        raise HTTPException(404, "Файл версии документа не найден")
    return metadata, target


def public_document_version(document_id: int, metadata: dict[str, Any], path: Path) -> dict[str, Any]:
    stat = path.stat()
    version_id = str(metadata.get("id") or f"v{metadata.get('version', '')}")
    return {
        "id": version_id,
        "version": metadata.get("version"),
        "name": str(metadata.get("name") or path.name),
        "size": stat.st_size,
        "file_extension": path.suffix.casefold(),
        "media_type": media_type_for_path(path),
        "created_at": str(metadata.get("created_at") or datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()),
        "download_url": f"/api/documents/{document_id}/versions/{version_id}",
    }


def current_object(db) -> ObjectRecord:
    obj = db.scalar(select(ObjectRecord).where(ObjectRecord.is_current.is_(True)))
    if not obj:
        obj = db.scalar(select(ObjectRecord).where(ObjectRecord.status == "active").limit(1))
    if not obj:
        raise HTTPException(404, "Нет активного объекта")
    return obj


def serialize(model) -> dict:
    data = {column.name: getattr(model, column.name) for column in model.__table__.columns}
    for key, value in tuple(data.items()):
        if hasattr(value, "isoformat"):
            data[key] = value.isoformat()
    return data


def document_primary_file(document: Document) -> Path | None:
    return safe_vault_file(document.source_file) or safe_vault_file(document.output_file)


def document_file_info(document: Document) -> dict[str, Any] | None:
    path = document_primary_file(document)
    if not path:
        return None
    info = file_info_for_path(
        path,
        download_url=f"/api/source/{document.id}" if safe_vault_file(document.source_file) else f"/api/files/{document.id}",
        preview_url=f"/api/preview/{document.id}",
    )
    if info is not None:
        info["kind"] = "source" if safe_vault_file(document.source_file) else "generated"
    return info


def serialize_template_for_read(template: Template) -> dict[str, Any]:
    data = serialize(template)
    path = None
    if template.template_file:
        try:
            path = resolve_template_path(template.template_file)
        except HTTPException:
            path = None
    data["template_file"] = path.name if path else ""
    data["file_info"] = file_info_for_path(
        path,
        download_url=f"/api/templates/{template.id}/file",
        preview_url=f"/api/templates/{template.id}/preview",
        allowed_roots=template_roots(),
    )
    return data


def document_has_file(document: Document) -> bool:
    return document_primary_file(document) is not None


def document_is_ready(document: Document) -> bool:
    return document.completion_status == "completed" and document_has_file(document)


def get_nested(data: dict, path: str):
    current = data
    for part in path.split("."):
        if not isinstance(current, dict) or part not in current:
            return None
        current = current[part]
    return current


def integrity_for_object(db, obj: ObjectRecord) -> dict:
    documents = db.scalars(select(Document).where(Document.object_id == obj.id)).all()
    all_rules = db.scalars(select(StructureRule).where(StructureRule.profile == obj.profile).order_by(StructureRule.id)).all()
    groups = db.scalars(select(ObjectGroup).where(ObjectGroup.object_id == obj.id)).all()
    groups_by_id = {group.id: group for group in groups}
    groups_by_kind: dict[str, list[ObjectGroup]] = {}
    for group in groups:
        groups_by_kind.setdefault(group.kind, []).append(group)
    rules = [
        rule
        for rule in all_rules
        if not groups_by_kind.get(rule.group_kind)
        or any(group.is_required for group in groups_by_kind[rule.group_kind])
    ]
    exceptions = {
        item.rule_id: item
        for item in db.scalars(select(StructureException).where(StructureException.object_id == obj.id, StructureException.status == "accepted")).all()
    }
    raw_group_exceptions = (obj.meta_json or {}).get("group_integrity_exceptions", {})
    group_exceptions = raw_group_exceptions if isinstance(raw_group_exceptions, dict) else {}
    checks = []
    satisfied = 0
    for rule in rules:
        required_group_ids = {
            group.id
            for group in groups_by_kind.get(rule.group_kind, [])
            if group.is_required
        }
        matches = [
            document
            for document in documents
            if document.doc_type in rule.doc_types_json
            and document.status != "superseded"
            and document.group_id in required_group_ids
            and groups_by_id.get(document.group_id, None) is not None
        ]
        present_matches = [document for document in matches if document_has_file(document)]
        completed_matches = [document for document in present_matches if document.completion_status == "completed"]
        records_count = len(matches)
        present_count = len(present_matches)
        exception = exceptions.get(rule.id)
        if present_count >= rule.min_count:
            status = "ok"
            satisfied += 1
        elif exception and exception.resolution_type != "pending":
            status = "waived"
            satisfied += 1
        elif exception:
            status = "pending"
        elif records_count >= rule.min_count:
            status = "pending"
        else:
            status = rule.severity
        checks.append(
            {
                "id": rule.id,
                "code": rule.code,
                "title": rule.title,
                "group_kind": rule.group_kind,
                "doc_types": rule.doc_types_json,
                "description": rule.description,
                "required": rule.min_count,
                "found": present_count,
                "records": records_count,
                "present_files": present_count,
                "completed_files": len(completed_matches),
                "ready": len(completed_matches),
                "document_ids": [document.id for document in matches],
                "pending_document_id": next(
                    (
                        document.id
                        for document in matches
                        if not document_has_file(document)
                    ),
                    None,
                ),
                "status": status,
                "reason": exception.reason if exception else "",
                "resolution_type": exception.resolution_type if exception else "",
                "evidence": exception.evidence if exception else "",
                "resolved_by": exception.created_by if exception else "",
                "resolved_at": exception.updated_at.isoformat() if exception and exception.updated_at else "",
            }
        )
    custom_required_groups = sorted(
        (
            group
            for group in groups
            if group.is_required and group.code not in STANDARD_GROUP_CODES
        ),
        key=lambda group: (group.sort_order, group.id),
    )
    for group in custom_required_groups:
        matches = [
            document
            for document in documents
            if document.group_id == group.id and document.status != "superseded"
        ]
        present_matches = [document for document in matches if document_has_file(document)]
        completed_matches = [document for document in present_matches if document.completion_status == "completed"]
        exception = group_exceptions.get(str(group.id), {})
        accepted_exception = exception if isinstance(exception, dict) and exception.get("status", "accepted") == "accepted" else {}
        resolution_type = str(accepted_exception.get("resolution_type") or "")
        if present_matches:
            status = "ok"
            satisfied += 1
        elif accepted_exception and resolution_type != "pending":
            status = "waived"
            satisfied += 1
        elif accepted_exception or matches:
            status = "pending"
        else:
            status = "error"
        checks.append(
            {
                "id": -group.id,
                "code": f"required-group-{group.id}",
                "title": f"Обязательная группа: {group.title}",
                "group_kind": group.kind,
                "doc_types": ["*"],
                "description": group.description or "В обязательной пользовательской группе должен быть хотя бы один файл.",
                "required": 1,
                "found": len(present_matches),
                "records": len(matches),
                "present_files": len(present_matches),
                "completed_files": len(completed_matches),
                "ready": len(completed_matches),
                "document_ids": [document.id for document in matches],
                "pending_document_id": next((document.id for document in matches if not document_has_file(document)), None),
                "status": status,
                "reason": str(accepted_exception.get("reason") or ""),
                "resolution_type": resolution_type,
                "evidence": str(accepted_exception.get("evidence") or ""),
                "resolved_by": str(accepted_exception.get("created_by") or ""),
                "resolved_at": str(accepted_exception.get("updated_at") or ""),
            }
        )
    return {
        "score": round(satisfied / len(checks) * 100) if checks else 100,
        "total": len(checks),
        "satisfied": satisfied,
        "issues": sum(check["status"] in {"error", "warning", "pending"} for check in checks),
        "checks": checks,
    }


def resolved_template_fields(template: Template, obj: ObjectRecord, preset: Preset | None = None, documents: list[Document] | None = None) -> dict:
    meta = obj.meta_json or {}
    object_data = {
        "object": {
            "name": obj.name,
            "address": obj.address,
            "code": obj.code,
            "meta": meta,
            "participants": [
                value
                for value in (
                    meta.get("customer", {}).get("name"),
                    meta.get("contractor", {}).get("name"),
                    meta.get("responsible_person", {}).get("fio"),
                )
                if value
            ],
            "documents": [
                {"number": document.number, "title": document.title, "date": document.document_date, "status": document.status}
                for document in (documents or [])
            ],
        }
    }
    preset_data = {"preset": preset.values_json if preset else {}}
    resolved = {}
    fields = []
    for field in template.fields_json or []:
        source = field.get("source") or template.field_map_json.get(field["key"], "manual")
        value = None
        if source.startswith("object."):
            value = get_nested(object_data, source)
            if value in (None, "", []) and source == "object.meta.construction_type":
                value = meta.get("work_kind")
        elif source.startswith("preset."):
            value = get_nested(preset_data, source)
        elif source == "default":
            value = field.get("default", "Соответствует")
        if value in (None, "", []) and preset and field["key"] in (preset.values_json or {}):
            value = preset.values_json[field["key"]]
        status = "resolved" if value not in (None, "", []) else "manual"
        fields.append({**field, "source": source, "value": value, "status": status})
        if status == "resolved":
            resolved[field["key"]] = value
    return {
        "template": serialize_template_for_read(template),
        "object": serialize(obj),
        "preset": serialize(preset) if preset else None,
        "fields": fields,
        "values": resolved,
        "resolved_count": sum(field["status"] == "resolved" for field in fields),
        "manual_count": sum(field["status"] == "manual" for field in fields),
    }


def resolved_document_values(
    document: Document,
    template: Template | None,
    obj: ObjectRecord | None,
    documents: list[Document] | None = None,
) -> dict:
    """Return document values with object-bound fields refreshed from the object profile."""
    stored_values = dict(document.values_json or {})
    if not template or not obj:
        return stored_values

    resolution = resolved_template_fields(template, obj, documents=documents)
    values = {**resolution["values"], **stored_values}
    for field in resolution["fields"]:
        if not str(field.get("source") or "").startswith("object."):
            continue
        key = field["key"]
        if field["status"] == "resolved":
            values[key] = field["value"]
        else:
            # An empty current profile value must not reveal an obsolete snapshot.
            values.pop(key, None)
    return values


def serialize_documents_for_read(db, documents: list[Document]) -> list[dict]:
    """Serialize documents using current object data without mutating the database."""
    object_map = {object_id: db.get(ObjectRecord, object_id) for object_id in {item.object_id for item in documents}}
    template_map = {
        template_id: db.get(Template, template_id)
        for template_id in {item.template_id for item in documents if item.template_id}
    }
    documents_by_object: dict[int, list[Document]] = {}
    for item in documents:
        documents_by_object.setdefault(item.object_id, []).append(item)

    result = []
    for item in documents:
        data = serialize(item)
        source = safe_vault_file(item.source_file)
        output = safe_vault_file(item.output_file)
        data["source_file"] = source.name if source else ""
        data["output_file"] = output.name if output else None
        data["file_info"] = document_file_info(item)
        values = resolved_document_values(
            item,
            template_map.get(item.template_id),
            object_map.get(item.object_id),
            documents_by_object.get(item.object_id),
        )
        values.pop("_previous_output_file", None)
        data["values_json"] = values
        data["missing_fields_json"] = missing_fields_for_template(
            template_map.get(item.template_id),
            values,
            item.number,
            item.work_title,
            item.document_date,
        )
        result.append(data)
    return result


def missing_fields_for_template(template: Template | None, values: dict, number: str = "", work_title: str = "", document_date: str = "") -> list[str]:
    if not template:
        return []
    aliases = {
        "act.number": number,
        "order.number": number,
        "scheme.number": number,
        "journal.number": number,
        "act.date": document_date,
        "order.date": document_date,
        "works": work_title,
        "scope": work_title,
    }
    missing = []
    for field in template.fields_json or []:
        if not field.get("required"):
            continue
        key = field["key"]
        value = values.get(key)
        if value in (None, "", []):
            value = get_nested(values, key)
        if value in (None, "", []):
            value = aliases.get(key)
        if value in (None, "", []):
            missing.append(key)
    return missing


@router.get("/app-state")
def app_state() -> dict:
    with SessionLocal() as db:
        obj = current_object(db)
        objects = db.scalars(select(ObjectRecord).order_by(ObjectRecord.created_at)).all()
        sections = db.scalars(select(Section).order_by(Section.sort_order)).all()
        documents = db.scalars(select(Document).where(Document.object_id == obj.id).order_by(Document.id.desc())).all()
        return {
            "ok": True,
            "appTitle": "ID Vault",
            "activeObjectId": obj.id,
            "currentObject": serialize(obj),
            "objects": [serialize(item) for item in objects if item.status == "active"],
            "archiveObjects": [serialize(item) for item in objects if item.status == "archived"],
            "sections": [serialize(item) for item in sections],
            "documents": serialize_documents_for_read(db, documents),
            "capabilities": {"pdf": False, "fts": True, "offline": True},
        }


@router.get("/objects/summary")
def object_summaries() -> list[dict]:
    with SessionLocal() as db:
        objects = db.scalars(select(ObjectRecord).order_by(ObjectRecord.is_current.desc(), ObjectRecord.updated_at.desc())).all()
        result = []
        for obj in objects:
            documents = db.scalars(select(Document).where(Document.object_id == obj.id)).all()
            integrity = integrity_for_object(db, obj)
            result.append(
                {
                    **serialize(obj),
                    "documents_count": len(documents),
                    "records": len(documents),
                    "present_files": sum(document_has_file(document) for document in documents),
                    "completed_files": sum(document_is_ready(document) for document in documents),
                    "ready_count": sum(document_is_ready(document) for document in documents),
                    "issues_count": integrity["issues"] + sum(len(document.warnings_json) for document in documents),
                    "integrity": integrity["score"],
                    "updated_label": "Сегодня" if obj.is_current else "Недавно",
                }
            )
        return result


@router.get("/objects/{object_id}/workspace")
def object_workspace(object_id: int) -> dict:
    with SessionLocal() as db:
        obj = db.get(ObjectRecord, object_id)
        if not obj:
            raise HTTPException(404, "Объект не найден")
        groups = db.scalars(select(ObjectGroup).where(ObjectGroup.object_id == obj.id).order_by(ObjectGroup.sort_order)).all()
        # Keep registry order stable. Editing or importing a later file must not
        # reshuffle an established object register.
        documents = db.scalars(select(Document).where(Document.object_id == obj.id).order_by(Document.created_at, Document.id)).all()
        document_map: dict[int | None, list[Document]] = {}
        for document in documents:
            document_map.setdefault(document.group_id, []).append(document)
        serialized_documents = serialize_documents_for_read(db, documents)
        return {
            "object": serialize(obj),
            "records": len(documents),
            "present_files": sum(document_has_file(document) for document in documents),
            "completed_files": sum(document_is_ready(document) for document in documents),
            "groups": [
                {
                    **serialize(group),
                    "documents_count": len(document_map.get(group.id, [])),
                    "records": len(document_map.get(group.id, [])),
                    "present_files": sum(document_has_file(document) for document in document_map.get(group.id, [])),
                    "completed_files": sum(document_is_ready(document) for document in document_map.get(group.id, [])),
                    "ready_count": sum(document_is_ready(document) for document in document_map.get(group.id, [])),
                    "issues_count": sum(bool(document.warnings_json or document.missing_fields_json) for document in document_map.get(group.id, [])),
                }
                for group in groups
            ],
            "documents": serialized_documents,
            "integrity": integrity_for_object(db, obj),
        }


@router.get("/objects", response_model=list[ObjectOut])
def list_objects(include_archived: bool = False):
    with SessionLocal() as db:
        query = select(ObjectRecord).order_by(ObjectRecord.created_at.desc())
        if not include_archived:
            query = query.where(ObjectRecord.status == "active")
        return db.scalars(query).all()


@router.post("/object/select")
def select_object(payload: dict) -> dict:
    with SessionLocal.begin() as db:
        obj = db.get(ObjectRecord, int(payload.get("id", 0)))
        if not obj:
            raise HTTPException(404, "Объект не найден")
        for item in db.scalars(select(ObjectRecord).where(ObjectRecord.is_current.is_(True))):
            item.is_current = False
        obj.is_current = True
    return app_state()


@router.post("/object/add", response_model=ObjectOut)
def add_object(payload: ObjectCreate):
    with SessionLocal.begin() as db:
        next_id = (db.scalar(select(ObjectRecord.id).order_by(ObjectRecord.id.desc()).limit(1)) or 0) + 1
        code = payload.code or f"{date.today().year}-{next_id:03d}"
        if db.scalar(select(ObjectRecord.id).where(ObjectRecord.code == code)):
            raise HTTPException(409, "Код объекта уже используется")
        for item in db.scalars(select(ObjectRecord).where(ObjectRecord.is_current.is_(True))):
            item.is_current = False
        obj = ObjectRecord(
            code=code,
            name=payload.name,
            short_name=payload.short_name or payload.name[:80],
            address=payload.address,
            comment=payload.comment,
            folder_path=f"objects/{safe_component(code, 'OBJECT')}_{safe_component(payload.name[:40], 'Объект')}",
            is_current=True,
            profile="industrial_repair",
            meta_json={},
        )
        db.add(obj)
        db.flush()
        for code_value, title, kind, description, order, required in GROUPS:
            db.add(ObjectGroup(object_id=obj.id, code=code_value, title=title, kind=kind, description=description, sort_order=order, is_required=required))
        return obj


@router.post("/objects/{object_id}/groups")
def add_group(object_id: int, payload: dict) -> dict:
    title = str(payload.get("title") or "").strip()
    if len(title) < 2:
        raise HTTPException(422, "Укажите название группы")
    with SessionLocal.begin() as db:
        obj = db.get(ObjectRecord, object_id)
        if not obj:
            raise HTTPException(404, "Объект не найден")
        existing = db.scalar(select(ObjectGroup).where(ObjectGroup.object_id == object_id, ObjectGroup.title == title))
        if existing:
            raise HTTPException(409, "Такая группа уже существует")
        last_order = db.scalar(select(ObjectGroup.sort_order).where(ObjectGroup.object_id == object_id).order_by(ObjectGroup.sort_order.desc()).limit(1)) or 0
        code = f"custom-{object_id}-{last_order + 10}"
        group = ObjectGroup(
            object_id=object_id,
            code=code,
            title=title,
            kind=str(payload.get("kind") or "other"),
            description=str(payload.get("description") or ""),
            sort_order=last_order + 10,
            is_required=bool(payload.get("is_required", False)),
        )
        db.add(group)
        db.flush()
        return serialize(group)


@router.patch("/objects/{object_id}", response_model=ObjectOut)
def update_object(object_id: int, payload: ObjectPatch):
    with SessionLocal.begin() as db:
        obj = db.get(ObjectRecord, object_id)
        if not obj:
            raise HTTPException(404, "Объект не найден")
        for key, value in payload.model_dump(exclude_none=True).items():
            setattr(obj, key, value)
        db.flush()
        return obj


@router.post("/object/archive")
def archive_object(payload: dict) -> dict:
    with SessionLocal.begin() as db:
        obj = db.get(ObjectRecord, int(payload.get("id", 0)))
        if not obj:
            raise HTTPException(404, "Объект не найден")
        replacement = db.scalar(select(ObjectRecord).where(ObjectRecord.status == "active", ObjectRecord.id != obj.id).limit(1))
        if not replacement:
            raise HTTPException(409, "Нельзя архивировать последний активный объект")
        obj.status = "archived"
        obj.is_current = False
        replacement.is_current = True
    return app_state()


@router.post("/object/unarchive")
def unarchive_object(payload: dict) -> dict:
    with SessionLocal.begin() as db:
        obj = db.get(ObjectRecord, int(payload.get("id", 0)))
        if not obj:
            raise HTTPException(404, "Объект не найден")
        obj.status = "active"
    return app_state()


@router.get("/documents", response_model=list[DocumentOut])
def list_documents(object_id: int | None = None, q: str = ""):
    with SessionLocal() as db:
        oid = object_id or current_object(db).id
        query = select(Document).where(Document.object_id == oid)
        if q:
            query = query.where(or_(Document.number.ilike(f"%{q}%"), Document.work_title.ilike(f"%{q}%")))
        documents = db.scalars(query.order_by(Document.created_at, Document.id)).all()
        return serialize_documents_for_read(db, documents)


@router.post("/documents", response_model=DocumentOut)
def create_document(payload: DocumentCreate):
    with SessionLocal.begin() as db:
        obj = db.get(ObjectRecord, payload.object_id) if payload.object_id else current_object(db)
        if not obj:
            raise HTTPException(404, "Объект не найден")
        values = dict(payload.values_json)
        template = db.get(Template, payload.template_id) if payload.template_id else None
        if payload.template_id and not template:
            raise HTTPException(404, "Шаблон не найден")
        missing_fields: list[str] = []
        if payload.preset_id:
            preset = db.get(Preset, payload.preset_id)
            if not preset:
                raise HTTPException(404, "Пресет не найден")
            if template and preset.doc_type != template.doc_type:
                raise HTTPException(422, "Пресет не подходит к выбранному шаблону")
            values = {**preset.values_json, **values}
            preset.use_count += 1
        else:
            preset = None
        if template:
            object_documents = db.scalars(select(Document).where(Document.object_id == obj.id)).all()
            resolution = resolved_template_fields(template, obj, preset, object_documents)
            values = {**resolution["values"], **values}
        group_id = payload.group_id
        if group_id:
            selected_group = db.get(ObjectGroup, group_id)
            if not selected_group or selected_group.object_id != obj.id:
                raise HTTPException(422, "Группа не принадлежит выбранному объекту")
        if not group_id and template:
            group = db.scalar(select(ObjectGroup).where(ObjectGroup.object_id == obj.id, ObjectGroup.kind == template.group_kind).limit(1))
            group_id = group.id if group else None
        if not group_id and not template:
            group = auto_object_group(db, obj, payload.doc_type, payload.title)
            group_id = group.id if group else None
        count = db.scalar(select(Document.id).where(Document.object_id == obj.id).order_by(Document.id.desc()).limit(1)) or 0
        number = payload.number or f"ИД-{count + 15:03d}"
        document_date = payload.document_date or date.today().strftime("%d.%m.%Y")
        missing_fields = missing_fields_for_template(template, values, number, payload.work_title, document_date)
        doc = Document(
            object_id=obj.id,
            group_id=group_id,
            template_id=template.id if template else None,
            number=number,
            title=payload.title,
            doc_type=template.doc_type if template else payload.doc_type,
            work_title=payload.work_title or values.get("work_title", ""),
            document_date=document_date,
            status="draft",
            completion_status="not_started",
            values_json=values,
            warnings_json=[],
            missing_fields_json=missing_fields,
        )
        db.add(doc)
        db.flush()
        return DocumentOut.model_validate(serialize_documents_for_read(db, [doc])[0])


@router.post("/objects/{object_id}/import", response_model=DocumentOut)
def import_document_file(object_id: int, file: UploadFile = File(...), group_id: int | None = Form(default=None)):
    with SessionLocal.begin() as db:
        obj = db.get(ObjectRecord, object_id)
        if not obj:
            raise HTTPException(404, "Объект не найден")
        source_name = Path(file.filename or "Документ").name
        if Path(source_name).suffix.casefold() not in DOCUMENT_SUFFIXES:
            raise HTTPException(415, "Поддерживаются документы, таблицы, PDF, изображения и чертежи")
        if getattr(file, "size", 0) and file.size > MAX_DOCUMENT_SIZE:
            raise HTTPException(413, "Файл больше 100 МБ")
        doc_type = infer_document_type(source_name)
        group = db.get(ObjectGroup, group_id) if group_id else None
        if group and group.object_id != obj.id:
            raise HTTPException(422, "Группа не принадлежит выбранному объекту")
        if not group:
            kind_by_type = {
                "registry": "registry",
                "general_journal": "journals",
                "incoming_control": "quality",
                "welding_journal": "journals",
                "aosr": "acts",
                "vik": "acts",
                "hydro": "acts",
                "executive_scheme": "schemes",
                "certificate": "quality",
                "passport": "quality",
                "appointment_order": "permits",
                "ppr": "source",
            }
            kind = kind_by_type.get(doc_type, "other")
            group = db.scalar(select(ObjectGroup).where(ObjectGroup.object_id == obj.id, ObjectGroup.kind == kind).limit(1))
        filename = safe_file_name(source_name)
        object_root = object_storage_root(obj)
        target_dir = object_root / "Импорт" / safe_component(group.title if group else "Без группы")
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / filename
        suffix = 2
        while target.exists():
            target = target_dir / f"{Path(filename).stem}_{suffix}{Path(filename).suffix}"
            suffix += 1
        temporary = target.with_name(f".{target.stem}.{uuid4().hex}.uploading{target.suffix}")
        write_upload_limited(file, temporary, MAX_DOCUMENT_SIZE, "Файл больше 100 МБ")
        os.replace(temporary, target)
        next_id = (db.scalar(select(Document.id).order_by(Document.id.desc()).limit(1)) or 0) + 1
        doc = Document(
            object_id=obj.id,
            group_id=group.id if group else None,
            number=f"ИМП-{next_id:04d}",
            title=Path(source_name).stem,
            doc_type=doc_type,
            work_title="Импортированный документ",
            document_date=date.today().strftime("%d.%m.%Y"),
            status="ready",
            completion_status="completed",
            values_json={},
            warnings_json=[],
            missing_fields_json=[],
            source_kind="imported",
            source_file=str(target),
        )
        db.add(doc)
        db.flush()
        return DocumentOut.model_validate(serialize_documents_for_read(db, [doc])[0])


@router.post("/documents/{document_id}/replace", response_model=DocumentOut, include_in_schema=False)
@router.post("/documents/{document_id}/file", response_model=DocumentOut)
def replace_document_file(document_id: int, file: UploadFile = File(...)):
    source_name = Path(file.filename or "Документ").name
    if Path(source_name).suffix.casefold() not in DOCUMENT_SUFFIXES:
        raise HTTPException(415, "Поддерживаются документы, таблицы, PDF, изображения и чертежи")
    if getattr(file, "size", 0) and file.size > MAX_DOCUMENT_SIZE:
        raise HTTPException(413, "Файл больше 100 МБ")

    db = SessionLocal()
    temporary: Path | None = None
    target: Path | None = None
    old_current: Path | None = None
    archived: Path | None = None
    archived_was_move = False
    try:
        document = db.get(Document, document_id)
        if not document:
            raise HTTPException(404, "Документ не найден")
        obj = db.get(ObjectRecord, document.object_id)
        if not obj:
            raise HTTPException(404, "Объект документа не найден")
        object_root = object_storage_root(obj)
        old_current = document_primary_file(document)
        versions_root = (object_root / "versions").resolve()
        if old_current and is_within(old_current, object_root) and not is_within(old_current, versions_root):
            target_dir = old_current.parent
        else:
            group = db.get(ObjectGroup, document.group_id) if document.group_id else None
            target_dir = object_root / "Документы" / safe_component(group.title if group else "Без группы")
        target_dir.mkdir(parents=True, exist_ok=True)
        temporary = target_dir / f".{safe_component(source_name)}.{uuid4().hex}.uploading"
        write_upload_limited(file, temporary, MAX_DOCUMENT_SIZE, "Файл больше 100 МБ")

        metadata, archived, archived_was_move = archive_document_file(document, obj)
        target = unique_file_target(target_dir, source_name)
        os.replace(temporary, target)
        temporary = None

        values = dict(document.values_json or {})
        versions = document_version_entries(document)
        if metadata:
            versions.append(metadata)
            values["versions"] = versions
            values["file_version"] = int(metadata["version"]) + 1
        else:
            configured_version = values.get("file_version")
            values["file_version"] = int(configured_version) if str(configured_version or "").isdigit() else 1
            values.setdefault("versions", versions)
        document.values_json = values
        if document.output_file:
            document.values_json = {**document.values_json, "_previous_output_file": document.output_file}
            document.output_file = None
        document.source_file = str(target)
        document.source_kind = "uploaded"
        document.status = "ready"
        document.completion_status = "in_progress"
        db.flush()
        response = DocumentOut.model_validate(serialize_documents_for_read(db, [document])[0])
        db.commit()
        return response
    except Exception:
        db.rollback()
        if temporary:
            temporary.unlink(missing_ok=True)
        if target and target.exists():
            target.unlink(missing_ok=True)
        if archived and archived.exists():
            if archived_was_move and old_current:
                old_current.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(archived), str(old_current))
            elif not archived_was_move:
                archived.unlink(missing_ok=True)
        raise
    finally:
        db.close()


@router.get("/documents/{document_id}/versions")
def list_document_versions(document_id: int) -> dict[str, Any]:
    with SessionLocal() as db:
        document = db.get(Document, document_id)
        if not document:
            raise HTTPException(404, "Документ не найден")
        obj = db.get(ObjectRecord, document.object_id)
        if not obj:
            raise HTTPException(404, "Объект документа не найден")
        versions = []
        for metadata in document_version_entries(document):
            try:
                _metadata, path = resolve_document_version(document, obj, str(metadata.get("id") or metadata.get("version") or ""))
            except HTTPException:
                continue
            versions.append(public_document_version(document.id, metadata, path))
        return {
            "document_id": document.id,
            "current": document_file_info(document),
            "versions": sorted(versions, key=lambda item: int(item.get("version") or 0), reverse=True),
        }


@router.get("/documents/{document_id}/versions/{version_id}")
def download_document_version(document_id: int, version_id: str):
    with SessionLocal() as db:
        document = db.get(Document, document_id)
        if not document:
            raise HTTPException(404, "Документ не найден")
        obj = db.get(ObjectRecord, document.object_id)
        if not obj:
            raise HTTPException(404, "Объект документа не найден")
        metadata, path = resolve_document_version(document, obj, version_id)
        return FileResponse(
            path,
            filename=str(metadata.get("name") or path.name),
            media_type=media_type_for_path(path),
        )


@router.patch("/documents/{document_id}", response_model=DocumentOut)
def update_document(document_id: int, payload: DocumentPatch):
    with SessionLocal.begin() as db:
        doc = db.get(Document, document_id)
        if not doc:
            raise HTTPException(404, "Документ не найден")
        changes = payload.model_dump(exclude_none=True)
        values_patch = changes.pop("values_json", None)
        requested_completion = changes.pop("completion_status", None)
        previous_completion = doc.completion_status
        content_changed = any(
            key in changes and getattr(doc, key) != value
            for key, value in changes.items()
            if key in {"number", "work_title", "document_date"}
        )
        if values_patch is not None:
            merged_values = {**(doc.values_json or {}), **values_patch}
            content_changed = content_changed or merged_values != (doc.values_json or {})
            doc.values_json = merged_values
        for key, value in changes.items():
            setattr(doc, key, value)
        if content_changed and doc.output_file:
            doc.values_json = {**(doc.values_json or {}), "_previous_output_file": doc.output_file}
            doc.output_file = None
            doc.status = "draft"
        template = db.get(Template, doc.template_id) if doc.template_id else None
        if template:
            obj = db.get(ObjectRecord, doc.object_id)
            object_documents = db.scalars(select(Document).where(Document.object_id == doc.object_id)).all()
            doc.values_json = resolved_document_values(doc, template, obj, object_documents)
        doc.missing_fields_json = missing_fields_for_template(
            template,
            doc.values_json or {},
            doc.number,
            doc.work_title,
            doc.document_date,
        )
        if requested_completion is not None:
            doc.completion_status = requested_completion
        elif content_changed:
            doc.completion_status = "needs_revision" if previous_completion == "completed" else "in_progress"
        db.flush()
        return DocumentOut.model_validate(serialize_documents_for_read(db, [doc])[0])


@router.delete("/documents/{document_id}")
def delete_document(document_id: int) -> dict:
    with SessionLocal.begin() as db:
        doc = db.get(Document, document_id)
        if not doc:
            raise HTTPException(404, "Документ не найден")
        object_id = doc.object_id
        source_file = doc.source_file or ""
        output_candidates = [
            value
            for value in (
                doc.output_file,
                (doc.values_json or {}).get("_previous_output_file"),
            )
            if isinstance(value, str) and value
        ]
        db.delete(doc)

    deleted_outputs: list[str] = []
    preserved_outputs: list[str] = []
    deletion_errors: list[str] = []
    source_resolved = Path(source_file).expanduser().resolve() if source_file else None
    for value in dict.fromkeys(output_candidates):
        owned = safe_owned_output(value, document_id, object_id)
        if not owned or (source_resolved and owned == source_resolved):
            preserved_outputs.append(value)
            continue
        try:
            if owned.exists():
                owned.unlink()
                deleted_outputs.append(str(owned))
        except OSError as exc:
            deletion_errors.append(f"{owned}: {exc}")
    return {
        "ok": True,
        "document_id": document_id,
        "source_file_preserved": Path(source_file).name if source_file else "",
        "deleted_outputs": [Path(value).name for value in deleted_outputs],
        "preserved_outputs": [Path(value).name for value in preserved_outputs],
        "warnings": ["Не удалось удалить один или несколько сформированных файлов"] if deletion_errors else [],
    }


@router.post("/generate")
def generate_documents(payload: dict) -> dict:
    ids = payload.get("ids") or payload.get("nums") or []
    blocked: list[dict[str, Any]] = []
    made: list[dict[str, Any]] = []
    with SessionLocal.begin() as db:
        query = select(Document)
        if ids:
            numeric_ids = [int(value) for value in ids if str(value).isdigit()]
            query = query.where(or_(Document.id.in_(numeric_ids), Document.number.in_([str(value) for value in ids])))
        else:
            query = query.where(Document.object_id == current_object(db).id)
        docs = db.scalars(query).all()
        if ids and not docs:
            raise HTTPException(404, "Документы для формирования не найдены")

        plans: list[dict[str, Any]] = []
        for item in docs:
            obj = db.get(ObjectRecord, item.object_id)
            template = db.get(Template, item.template_id) if item.template_id else None
            if template:
                object_documents = db.scalars(select(Document).where(Document.object_id == item.object_id)).all()
                item.values_json = resolved_document_values(item, template, obj, object_documents)
            missing = missing_fields_for_template(
                template,
                item.values_json or {},
                item.number,
                item.work_title,
                item.document_date,
            )
            template_path: Path | None = None
            replacements: dict[str, Any] = {}
            marker_count = 0
            if template and template.template_file:
                template_path = resolve_template_path(template.template_file)
                if template_path.suffix.casefold() != ".docx":
                    raise HTTPException(422, f"Автоматическое формирование поддерживает только DOCX-шаблоны: «{template.title}»")
                try:
                    markers = inspect_docx_markers(template_path)
                except Exception as exc:
                    raise HTTPException(422, f"DOCX-шаблон «{template.title}» повреждён или не читается: {exc}") from exc
                replacements, marker_missing = resolve_marker_values(
                    markers,
                    template.marker_map or {},
                    document_marker_context(item, obj),
                )
                missing = sorted(set([*missing, *marker_missing]), key=str.casefold)
                marker_count = sum(marker["occurrences"] for marker in markers)
            item.missing_fields_json = missing
            if missing:
                item.completion_status = "needs_revision"
                blocked.append(
                    {
                        "id": item.id,
                        "number": item.number,
                        "missing_fields": missing,
                    }
                )
                continue
            plans.append(
                {
                    "document": item,
                    "template": template,
                    "template_path": template_path,
                    "replacements": replacements,
                    "marker_count": marker_count,
                }
            )

        if not blocked:
            for plan in plans:
                item: Document = plan["document"]
                target = expected_output_path(item.id, item.object_id)
                target.parent.mkdir(parents=True, exist_ok=True)
                temporary = target.with_name(f".{target.stem}.building.docx")
                try:
                    if plan["template_path"]:
                        replaced = render_docx_template(plan["template_path"], temporary, plan["replacements"])
                    else:
                        word = WordDocument()
                        word.add_heading(item.title, 0)
                        word.add_paragraph(f"№ {item.number} от {item.document_date}")
                        word.add_heading("Выполненные работы", level=1)
                        word.add_paragraph(item.work_title)
                        template = plan["template"]
                        fields = template.fields_json if template else [
                            {"key": "project_doc", "label": "Проектная документация"},
                            {"key": "materials", "label": "Материалы"},
                            {"key": "norms", "label": "Нормативные документы"},
                        ]
                        aliases = {
                            "act.number": item.number,
                            "order.number": item.number,
                            "scheme.number": item.number,
                            "act.date": item.document_date,
                            "order.date": item.document_date,
                            "works": item.work_title,
                            "scope": item.work_title,
                        }
                        for field in fields:
                            value = item.values_json.get(field["key"], aliases.get(field["key"], "Не указано"))
                            if isinstance(value, (dict, list)):
                                value = json.dumps(value, ensure_ascii=False, indent=2)
                            word.add_heading(field.get("label") or field["key"], level=2)
                            word.add_paragraph(str(value or "Не указано"))
                        word.save(temporary)
                        replaced = 0
                    os.replace(temporary, target)
                except Exception:
                    if temporary.exists():
                        temporary.unlink()
                    raise
                item.output_file = str(target)
                item.status = "generated"
                item.completion_status = "completed"
                item.missing_fields_json = []
                made.append(
                    {
                        "id": item.id,
                        "number": item.number,
                        "file": target.name,
                        "download_url": f"/api/files/{item.id}",
                        "template_file": Path(plan["template"].template_file).name if plan["template"] and plan["template"].template_file else "",
                        "markers_found": plan["marker_count"],
                        "markers_replaced": replaced,
                    }
                )

        warning_count = sum(len(item.warnings_json) for item in docs)

    if blocked:
        raise HTTPException(
            422,
            detail={
                "message": "Формирование заблокировано: заполнены не все значения",
                "missing_fields": blocked,
            },
        )
    return {"ok": True, "made": made, "warnings": warning_count}


@router.get("/templates")
def list_templates(include_archived: bool = False) -> list[dict]:
    with SessionLocal() as db:
        query = select(Template).order_by(Template.updated_at.desc())
        if not include_archived:
            query = query.where(Template.is_active.is_(True))
        return [serialize_template_for_read(item) for item in db.scalars(query).all()]


@router.post("/templates")
def create_template(
    file: UploadFile = File(...),
    title: str = Form(default=""),
    doc_type: str = Form(default="other"),
    group_kind: str = Form(default="other"),
    notes: str = Form(default=""),
    marker_map: str | None = Form(default=None),
) -> dict:
    source_name = Path(file.filename or "template.docx").name
    if Path(source_name).suffix.casefold() not in TEMPLATE_SUFFIXES:
        raise HTTPException(415, "Шаблон должен быть файлом DOCX, XLSX или PDF")
    parsed_map: dict[str, Any] = {}
    if marker_map is not None:
        try:
            parsed_map = normalized_marker_map(json.loads(marker_map))
        except json.JSONDecodeError as exc:
            raise HTTPException(422, "marker_map содержит некорректный JSON") from exc
    if getattr(file, "size", 0) and file.size > MAX_TEMPLATE_SIZE:
        raise HTTPException(413, "Файл шаблона больше 25 МБ")

    with SessionLocal.begin() as db:
        template = Template(
            title=title.strip() or Path(source_name).stem,
            doc_type=doc_type.strip() or "other",
            group_kind=group_kind.strip() or "other",
            version=1,
            is_default=False,
            is_active=True,
            fields_json=[],
            field_map_json={},
            marker_map=parsed_map,
            notes=notes.strip(),
        )
        db.add(template)
        db.flush()
        target, version = managed_template_target(template, source_name)
        temporary = target.with_name(f".{target.stem}.{uuid4().hex}.uploading{target.suffix}")
        write_upload_limited(file, temporary, MAX_TEMPLATE_SIZE, "Файл шаблона больше 25 МБ")
        try:
            validate_template_file(temporary)
            os.replace(temporary, target)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise
        template.template_file = stored_template_path(target)
        template.version = version
        db.flush()
        result = template_markers_payload(template)
        return {"template": serialize_template_for_read(template), **result}


@router.get("/templates/{template_id}/markers")
def list_template_markers(template_id: int) -> dict:
    with SessionLocal() as db:
        template = db.get(Template, template_id)
        if not template:
            raise HTTPException(404, "Шаблон не найден")
        try:
            return template_markers_payload(template)
        except Exception as exc:
            if isinstance(exc, HTTPException):
                raise
            raise HTTPException(422, f"DOCX-шаблон повреждён или не читается: {exc}") from exc


@router.post("/templates/{template_id}/file")
def upload_template_file(
    template_id: int,
    file: UploadFile = File(...),
    marker_map: str | None = Form(default=None),
) -> dict:
    source_name = Path(file.filename or "template.docx").name
    suffix = Path(source_name).suffix.casefold()
    if suffix not in TEMPLATE_SUFFIXES:
        raise HTTPException(415, "Шаблон должен быть файлом DOCX, XLSX или PDF")
    parsed_map: dict[str, Any] | None = None
    if marker_map is not None:
        try:
            parsed_map = normalized_marker_map(json.loads(marker_map))
        except json.JSONDecodeError as exc:
            raise HTTPException(422, "marker_map содержит некорректный JSON") from exc
    if getattr(file, "size", 0) and file.size > MAX_TEMPLATE_SIZE:
        raise HTTPException(413, "Файл шаблона больше 25 МБ")

    with SessionLocal.begin() as db:
        template = db.get(Template, template_id)
        if not template:
            raise HTTPException(404, "Шаблон не найден")
        target, version = managed_template_target(template, source_name)
        temporary = target.with_name(f".{target.stem}.{uuid4().hex}.uploading{target.suffix}")
        write_upload_limited(file, temporary, MAX_TEMPLATE_SIZE, "Файл шаблона больше 25 МБ")
        try:
            validate_template_file(temporary)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise
        os.replace(temporary, target)
        template.template_file = stored_template_path(target)
        template.version = version
        if parsed_map is not None:
            template.marker_map = parsed_map
        elif suffix != ".docx":
            template.marker_map = {}
        db.flush()
        result = template_markers_payload(template)
        return {"template": serialize_template_for_read(template), **result}


@router.post("/templates/{template_id}/bind")
def bind_template_file(template_id: int, payload: TemplateBind) -> dict:
    raw = Path(payload.template_file).expanduser()
    source = raw.resolve() if raw.is_absolute() else (BASE_DIR / raw).resolve()
    if source.suffix.casefold() not in TEMPLATE_SUFFIXES:
        raise HTTPException(415, "Шаблон должен быть файлом DOCX, XLSX или PDF")
    if not source.is_file():
        raise HTTPException(404, "Указанный файл шаблона не найден")
    validate_template_file(source)

    with SessionLocal.begin() as db:
        template = db.get(Template, template_id)
        if not template:
            raise HTTPException(404, "Шаблон не найден")
        if any(is_within(source, root) for root in template_roots()):
            target = source
            version = template.version + (1 if template.template_file and stored_template_path(source) != template.template_file else 0)
        else:
            target, version = managed_template_target(template, source.name)
            temporary = target.with_name(f".{target.stem}.{uuid4().hex}.binding{target.suffix}")
            shutil.copy2(source, temporary)
            os.replace(temporary, target)
        template.template_file = stored_template_path(target)
        template.version = version
        if payload.marker_map is not None:
            template.marker_map = normalized_marker_map(payload.marker_map)
        elif source.suffix.casefold() != ".docx":
            template.marker_map = {}
        db.flush()
        result = template_markers_payload(template)
        return {"template": serialize_template_for_read(template), **result}


@router.get("/templates/{template_id}/file")
def download_template_file(template_id: int):
    with SessionLocal() as db:
        template = db.get(Template, template_id)
        if not template:
            raise HTTPException(404, "Шаблон не найден")
        if not template.template_file:
            raise HTTPException(404, "Файл шаблона не загружен")
        path = resolve_template_path(template.template_file)
        return FileResponse(path, filename=path.name, media_type=media_type_for_path(path))


@router.get("/templates/{template_id}/preview")
def preview_template_file(template_id: int):
    with SessionLocal() as db:
        template = db.get(Template, template_id)
        if not template:
            raise HTTPException(404, "Шаблон не найден")
        if not template.template_file:
            raise HTTPException(404, "Файл шаблона не загружен")
        path = resolve_template_path(template.template_file)
        return FileResponse(
            path,
            filename=path.name,
            media_type=media_type_for_path(path),
            content_disposition_type="inline",
        )


@router.get("/templates/{template_id}/spreadsheet-preview")
def preview_template_spreadsheet(template_id: int) -> dict[str, Any]:
    with SessionLocal() as db:
        template = db.get(Template, template_id)
        if not template or not template.template_file:
            raise HTTPException(404, "Файл шаблона не найден")
        return spreadsheet_preview_data(resolve_template_path(template.template_file))


@router.patch("/templates/{template_id}")
def update_template(template_id: int, payload: TemplatePatch) -> dict:
    changes = payload.model_dump(exclude_none=True)
    for key in ("title", "doc_type", "group_kind"):
        if key in changes:
            changes[key] = str(changes[key]).strip()
            if not changes[key]:
                raise HTTPException(422, f"Поле {key} не может быть пустым")
    if "notes" in changes:
        changes["notes"] = str(changes["notes"]).strip()
    with SessionLocal.begin() as db:
        template = db.get(Template, template_id)
        if not template:
            raise HTTPException(404, "Шаблон не найден")
        for key, value in changes.items():
            setattr(template, key, value)
        db.flush()
        return serialize_template_for_read(template)


@router.delete("/templates/{template_id}")
def archive_template(template_id: int) -> dict:
    with SessionLocal.begin() as db:
        template = db.get(Template, template_id)
        if not template:
            raise HTTPException(404, "Шаблон не найден")
        template.is_active = False
        db.flush()
        return {"ok": True, "archived": True, "template": serialize_template_for_read(template)}


@router.post("/templates/{template_id}/duplicate")
def duplicate_template(template_id: int, payload: TemplateDuplicate | None = None) -> dict:
    payload = payload or TemplateDuplicate()
    with SessionLocal.begin() as db:
        source_template = db.get(Template, template_id)
        if not source_template:
            raise HTTPException(404, "Шаблон не найден")
        duplicate = Template(
            title=(payload.title or f"{source_template.title} — копия").strip(),
            doc_type=source_template.doc_type,
            group_kind=source_template.group_kind,
            version=1,
            is_default=False,
            is_active=True,
            fields_json=list(source_template.fields_json or []),
            field_map_json=dict(source_template.field_map_json or {}),
            marker_map=dict(source_template.marker_map or {}),
            source_example=source_template.source_example,
            notes=(payload.notes if payload.notes is not None else source_template.notes).strip(),
        )
        db.add(duplicate)
        db.flush()
        if source_template.template_file:
            source = resolve_template_path(source_template.template_file)
            target, version = managed_template_target(duplicate, source.name)
            temporary = target.with_name(f".{target.stem}.{uuid4().hex}.copying{target.suffix}")
            shutil.copy2(source, temporary)
            os.replace(temporary, target)
            duplicate.template_file = stored_template_path(target)
            duplicate.version = version
        db.flush()
        return serialize_template_for_read(duplicate)


@router.post("/templates/{template_id}/copy", response_model=DocumentOut)
def copy_template_to_object(template_id: int, payload: TemplateCopy):
    with SessionLocal.begin() as db:
        template = db.get(Template, template_id)
        if not template:
            raise HTTPException(404, "Шаблон не найден")
        if not template.is_active:
            raise HTTPException(409, "Шаблон находится в архиве")
        if not template.template_file:
            raise HTTPException(409, "У шаблона нет файла")
        source = resolve_template_path(template.template_file)
        obj = db.get(ObjectRecord, payload.object_id)
        if not obj:
            raise HTTPException(404, "Объект не найден")
        group = db.get(ObjectGroup, payload.group_id) if payload.group_id else None
        if group and group.object_id != obj.id:
            raise HTTPException(422, "Группа не принадлежит выбранному объекту")
        if not group:
            group = db.scalar(
                select(ObjectGroup)
                .where(ObjectGroup.object_id == obj.id, ObjectGroup.kind == template.group_kind)
                .order_by(ObjectGroup.sort_order)
                .limit(1)
            )
        target_dir = object_storage_root(obj) / "Из шаблонов" / safe_component(group.title if group else "Без группы")
        target_dir.mkdir(parents=True, exist_ok=True)
        target = unique_file_target(target_dir, source.name)
        temporary = target.with_name(f".{target.stem}.{uuid4().hex}.copying{target.suffix}")
        try:
            shutil.copy2(source, temporary)
            os.replace(temporary, target)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise
        next_id = (db.scalar(select(Document.id).order_by(Document.id.desc()).limit(1)) or 0) + 1
        document = Document(
            object_id=obj.id,
            group_id=group.id if group else None,
            template_id=template.id,
            number=(payload.number or f"ИД-{next_id:03d}").strip(),
            title=(payload.title or template.title).strip(),
            doc_type=template.doc_type,
            work_title="",
            document_date=date.today().strftime("%d.%m.%Y"),
            status="ready",
            completion_status="in_progress",
            values_json={"versions": [], "file_version": 1, "template_version": template.version},
            warnings_json=[],
            missing_fields_json=[],
            source_kind="template",
            source_file=str(target),
        )
        db.add(document)
        db.flush()
        return DocumentOut.model_validate(serialize_documents_for_read(db, [document])[0])


@router.get("/templates/{template_id}/resolve")
def resolve_template(template_id: int, object_id: int, preset_id: int | None = None) -> dict:
    with SessionLocal() as db:
        template = db.get(Template, template_id)
        obj = db.get(ObjectRecord, object_id)
        preset = db.get(Preset, preset_id) if preset_id else None
        if not template or not obj:
            raise HTTPException(404, "Шаблон или объект не найден")
        documents = db.scalars(select(Document).where(Document.object_id == obj.id)).all()
        return resolved_template_fields(template, obj, preset, documents)


@router.get("/presets")
def list_presets(q: str = "") -> list[dict]:
    with SessionLocal() as db:
        query = select(Preset)
        if q:
            query = query.where(or_(Preset.title.ilike(f"%{q}%"), Preset.description.ilike(f"%{q}%")))
        return [serialize(item) for item in db.scalars(query.order_by(Preset.is_favorite.desc(), Preset.use_count.desc())).all()]


@router.get("/catalogs/{catalog}")
def list_catalog(catalog: str) -> list[dict]:
    with SessionLocal() as db:
        return [serialize(item) for item in db.scalars(select(CatalogItem).where(CatalogItem.catalog == catalog).order_by(CatalogItem.title)).all()]


@router.get("/search")
def global_search(q: str = Query(min_length=2)) -> dict:
    needle = q.casefold()
    contains = lambda *values: any(needle in str(value or "").casefold() for value in values)
    with SessionLocal() as db:
        docs = [
            item
            for item in db.scalars(select(Document).order_by(Document.updated_at.desc())).all()
            if contains(item.number, item.title, item.work_title)
        ][:12]
        objects = [
            item
            for item in db.scalars(select(ObjectRecord).order_by(ObjectRecord.updated_at.desc())).all()
            if contains(item.code, item.name, item.short_name, item.address)
        ][:8]
        presets = [
            item
            for item in db.scalars(select(Preset).order_by(Preset.is_favorite.desc(), Preset.use_count.desc())).all()
            if contains(item.title, item.description)
        ][:8]
        catalogs = [
            item
            for item in db.scalars(select(CatalogItem).order_by(CatalogItem.title)).all()
            if contains(item.title, item.code, item.subtitle)
        ][:12]
        return {
            "objects": [serialize(x) for x in objects],
            "documents": serialize_documents_for_read(db, docs),
            "presets": [serialize(x) for x in presets],
            "catalogs": [serialize(x) for x in catalogs],
        }


@router.get("/reports/completeness")
def completeness_report(object_id: int | None = None) -> dict:
    with SessionLocal() as db:
        if object_id is not None:
            obj = db.get(ObjectRecord, object_id)
            if not obj:
                raise HTTPException(status_code=404, detail="Объект не найден")
        else:
            obj = current_object(db)
        docs = db.scalars(select(Document).where(Document.object_id == obj.id)).all()
        total = len(docs)
        ready = sum(document_is_ready(item) for item in docs)
        warnings = sum(len(item.warnings_json) for item in docs)
        integrity = integrity_for_object(db, obj)
        groups = db.scalars(select(ObjectGroup).where(ObjectGroup.object_id == obj.id).order_by(ObjectGroup.sort_order)).all()
        return {
            "object_id": obj.id,
            "progress": integrity["score"],
            "total": total,
            "records": total,
            "present_files": sum(document_has_file(item) for item in docs),
            "completed_files": ready,
            "ready": ready,
            "warnings": warnings,
            "sections": [
                {
                    "title": group.title,
                    "records": sum(document.group_id == group.id for document in docs),
                    "present_files": sum(document.group_id == group.id and document_has_file(document) for document in docs),
                    "completed_files": sum(document.group_id == group.id and document_is_ready(document) for document in docs),
                    "ready": sum(document.group_id == group.id and document_is_ready(document) for document in docs),
                    "total": sum(document.group_id == group.id for document in docs),
                    "status": (
                        "ok"
                        if any(document.group_id == group.id for document in docs)
                        and all(document_has_file(document) for document in docs if document.group_id == group.id)
                        else "warning"
                        if any(document.group_id == group.id for document in docs)
                        else "error"
                        if group.is_required
                        else "neutral"
                    ),
                }
                for group in groups
            ],
            "checks": integrity["checks"],
        }


@router.post("/objects/{object_id}/integrity/{rule_id}/waive")
def waive_integrity_issue(object_id: int, rule_id: int, payload: dict) -> dict:
    reason = str(payload.get("reason") or "").strip()
    resolution_type = str(payload.get("resolution_type") or "not_applicable").strip()
    evidence = str(payload.get("evidence") or "").strip()
    allowed_types = {"not_applicable", "equivalent", "justified_exception", "pending"}
    if resolution_type not in allowed_types:
        raise HTTPException(422, "Неизвестный тип решения")
    if len(reason) < 5:
        raise HTTPException(422, "Укажите причину исключения")
    if resolution_type == "equivalent" and len(evidence) < 3:
        raise HTTPException(422, "Для замены эквивалентом укажите документ-подтверждение")
    with SessionLocal.begin() as db:
        obj = db.get(ObjectRecord, object_id)
        if not obj:
            raise HTTPException(404, "Объект не найден")
        if rule_id < 0:
            group = db.get(ObjectGroup, -rule_id)
            if not group or group.object_id != object_id or not group.is_required or group.code in STANDARD_GROUP_CODES:
                raise HTTPException(404, "Обязательная группа не найдена")
            meta = dict(obj.meta_json or {})
            group_exceptions = dict(meta.get("group_integrity_exceptions") or {})
            group_exceptions[str(group.id)] = {
                "resolution_type": resolution_type,
                "reason": reason,
                "evidence": evidence,
                "status": "accepted",
                "created_by": "Пользователь",
                "updated_at": datetime.now(timezone.utc).isoformat(),
            }
            meta["group_integrity_exceptions"] = group_exceptions
            obj.meta_json = meta
            db.flush()
            return integrity_for_object(db, obj)
        if not db.get(StructureRule, rule_id):
            raise HTTPException(404, "Правило не найдено")
        exception = db.scalar(select(StructureException).where(StructureException.object_id == object_id, StructureException.rule_id == rule_id))
        if exception:
            exception.resolution_type = resolution_type
            exception.reason = reason
            exception.evidence = evidence
            exception.status = "accepted"
        else:
            db.add(
                StructureException(
                    object_id=object_id,
                    rule_id=rule_id,
                    resolution_type=resolution_type,
                    reason=reason,
                    evidence=evidence,
                )
            )
    with SessionLocal() as db:
        return integrity_for_object(db, db.get(ObjectRecord, object_id))


@router.post("/objects/{object_id}/integrity/{rule_id}/restore")
def restore_integrity_issue(object_id: int, rule_id: int) -> dict:
    with SessionLocal.begin() as db:
        obj = db.get(ObjectRecord, object_id)
        if not obj:
            raise HTTPException(404, "Объект не найден")
        if rule_id < 0:
            group = db.get(ObjectGroup, -rule_id)
            if not group or group.object_id != object_id or group.code in STANDARD_GROUP_CODES:
                raise HTTPException(404, "Обязательная группа не найдена")
            meta = dict(obj.meta_json or {})
            group_exceptions = dict(meta.get("group_integrity_exceptions") or {})
            group_exceptions.pop(str(group.id), None)
            meta["group_integrity_exceptions"] = group_exceptions
            obj.meta_json = meta
            db.flush()
            return integrity_for_object(db, obj)
        db.execute(delete(StructureException).where(StructureException.object_id == object_id, StructureException.rule_id == rule_id))
    with SessionLocal() as db:
        obj = db.get(ObjectRecord, object_id)
        if not obj:
            raise HTTPException(404, "Объект не найден")
        return integrity_for_object(db, obj)


@router.get("/corpus/insights")
def corpus_insights() -> dict:
    with SessionLocal() as db:
        sources = db.scalars(select(CorpusSource).order_by(CorpusSource.updated_at.desc())).all()
        templates = db.scalars(select(Template).order_by(Template.group_kind, Template.title)).all()
        return {
            "sources": [serialize(source) for source in sources],
            "template_families": [
                {
                    "id": template.id,
                    "title": template.title,
                    "doc_type": template.doc_type,
                    "group_kind": template.group_kind,
                    "fields_count": len(template.fields_json or []),
                    "source_example": template.source_example,
                }
                for template in templates
            ],
            "principle": "Данные объекта подставляются автоматически; неизвестные значения остаются явными полями для заполнения.",
            "legal_note": "Публичные образцы используются как справочный корпус, а не как нормативная истина. Правила объекта должны подтверждаться актуальными требованиями договора, проекта и законодательства.",
        }


@router.post("/open-object-folder")
def open_object_folder(payload: dict) -> dict:
    with SessionLocal() as db:
        obj = db.get(ObjectRecord, int(payload.get("id", 0))) or current_object(db)
        target = (settings.vault_dir / obj.folder_path).resolve()
        target.mkdir(parents=True, exist_ok=True)
        os.startfile(target)
        return {"ok": True, "path": obj.folder_path, "folder_name": target.name}


@router.get("/files/{document_id}")
def download_document(document_id: int):
    with SessionLocal() as db:
        doc = db.get(Document, document_id)
        output = safe_vault_file(doc.output_file) if doc else None
        if not output:
            raise HTTPException(404, "Сформированный файл не найден")
        return FileResponse(output, filename=output.name, media_type=media_type_for_path(output))


@router.get("/source/{document_id}")
def open_source_document(document_id: int):
    with SessionLocal() as db:
        doc = db.get(Document, document_id)
        source = safe_vault_file(doc.source_file) if doc else None
        if not source:
            raise HTTPException(404, "Исходный файл не найден")
        return FileResponse(source, filename=source.name, media_type=media_type_for_path(source))


@router.get("/preview/{document_id}")
def preview_document(document_id: int):
    with SessionLocal() as db:
        doc = db.get(Document, document_id)
        path = document_primary_file(doc) if doc else None
        if not path:
            raise HTTPException(404, "Файл документа не найден")
        return FileResponse(
            path,
            filename=path.name,
            media_type=media_type_for_path(path),
            content_disposition_type="inline",
        )


@router.get("/preview/{document_id}/spreadsheet")
def preview_document_spreadsheet(document_id: int) -> dict[str, Any]:
    with SessionLocal() as db:
        document = db.get(Document, document_id)
        path = document_primary_file(document) if document else None
        if not path:
            raise HTTPException(404, "Файл документа не найден")
        return spreadsheet_preview_data(path)
