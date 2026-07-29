from __future__ import annotations

import json
import re
import shutil
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from docx.document import Document as WordDocumentType
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


MARKER_PATTERN = re.compile(
    r"(?P<token>"
    r"\[\[\s*(?P<square>[\w.:-]+)\s*\]\]"
    r"|\{\{\s*(?P<braces>[\w.:-]+)\s*\}\}"
    r"|<<\s*(?P<angles>[\w.:-]+)\s*>>"
    r")",
    re.UNICODE,
)


def marker_name(match: re.Match[str]) -> str:
    return next(value for value in (match.group("square"), match.group("braces"), match.group("angles")) if value)


def _iter_table_paragraphs(table: Table, prefix: str, seen_cells: set[Any]) -> Iterator[tuple[str, Paragraph]]:
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell_element = cell._tc
            if cell_element in seen_cells:
                continue
            seen_cells.add(cell_element)
            cell_prefix = f"{prefix}.cell[{row_index},{cell_index}]"
            for paragraph_index, paragraph in enumerate(cell.paragraphs):
                yield f"{cell_prefix}.paragraph[{paragraph_index}]", paragraph
            for table_index, nested_table in enumerate(cell.tables):
                yield from _iter_table_paragraphs(nested_table, f"{cell_prefix}.table[{table_index}]", seen_cells)


def _iter_container_paragraphs(container: WordDocumentType | _Cell | Any, prefix: str) -> Iterator[tuple[str, Paragraph]]:
    seen_cells: set[Any] = set()
    for paragraph_index, paragraph in enumerate(container.paragraphs):
        yield f"{prefix}.paragraph[{paragraph_index}]", paragraph
    for table_index, table in enumerate(container.tables):
        yield from _iter_table_paragraphs(table, f"{prefix}.table[{table_index}]", seen_cells)


def iter_document_paragraphs(document: WordDocumentType) -> Iterator[tuple[str, Paragraph]]:
    """Yield body, table, header and footer paragraphs without rebuilding the document."""
    yield from _iter_container_paragraphs(document, "body")
    seen_parts: set[str] = set()
    for section_index, section in enumerate(document.sections):
        containers = (
            ("header", section.header),
            ("first_page_header", section.first_page_header),
            ("even_page_header", section.even_page_header),
            ("footer", section.footer),
            ("first_page_footer", section.first_page_footer),
            ("even_page_footer", section.even_page_footer),
        )
        for kind, container in containers:
            part_name = str(container.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            yield from _iter_container_paragraphs(container, f"section[{section_index}].{kind}")


def inspect_docx_markers(path: str | Path) -> list[dict[str, Any]]:
    document = WordDocument(str(path))
    found: dict[str, dict[str, Any]] = {}
    for location, paragraph in iter_document_paragraphs(document):
        for match in MARKER_PATTERN.finditer(paragraph.text):
            name = marker_name(match)
            item = found.setdefault(name, {"name": name, "tokens": [], "occurrences": 0, "locations": []})
            token = match.group("token")
            if token not in item["tokens"]:
                item["tokens"].append(token)
            item["occurrences"] += 1
            if location not in item["locations"]:
                item["locations"].append(location)
    return sorted(found.values(), key=lambda item: item["name"].casefold())


def _replace_span(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    runs = paragraph.runs
    if not runs:
        return
    cursor = 0
    start_run = start_offset = end_run = end_offset = None
    for index, run in enumerate(runs):
        next_cursor = cursor + len(run.text)
        if start_run is None and start < next_cursor:
            start_run, start_offset = index, start - cursor
        if start_run is None and start == next_cursor and index == len(runs) - 1:
            start_run, start_offset = index, len(run.text)
        if end_run is None and end <= next_cursor:
            end_run, end_offset = index, end - cursor
            break
        cursor = next_cursor
    if start_run is None or end_run is None or start_offset is None or end_offset is None:
        return
    first = runs[start_run]
    last = runs[end_run]
    prefix = first.text[:start_offset]
    suffix = last.text[end_offset:]
    if start_run == end_run:
        first.text = prefix + replacement + suffix
        return
    first.text = prefix + replacement
    for index in range(start_run + 1, end_run):
        runs[index].text = ""
    last.text = suffix


def replace_docx_markers(document: WordDocumentType, replacements: dict[str, Any]) -> int:
    """Replace marker spans in-place while retaining the formatting of their first run."""
    replacement_count = 0
    for _location, paragraph in iter_document_paragraphs(document):
        matches = list(MARKER_PATTERN.finditer(paragraph.text))
        for match in reversed(matches):
            name = marker_name(match)
            if name not in replacements:
                continue
            value = replacements[name]
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            _replace_span(paragraph, match.start(), match.end(), str(value))
            replacement_count += 1
    return replacement_count


def render_docx_template(source: str | Path, target: str | Path, replacements: dict[str, Any]) -> int:
    """Copy a DOCX first, then replace markers inside the copy."""
    source_path = Path(source)
    target_path = Path(target)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, target_path)
    document = WordDocument(str(target_path))
    count = replace_docx_markers(document, replacements)
    document.save(str(target_path))
    return count


def _lookup(context: dict[str, Any], source: str) -> tuple[bool, Any]:
    if source in context:
        return True, context[source]
    current: Any = context
    for part in source.split("."):
        if not isinstance(current, dict) or part not in current:
            return False, None
        current = current[part]
    return True, current


def resolve_marker_values(
    markers: list[dict[str, Any]],
    marker_map: dict[str, Any] | None,
    context: dict[str, Any],
) -> tuple[dict[str, Any], list[str]]:
    """Resolve marker names from document/object values and an optional marker map."""
    configured = marker_map or {}
    configured_casefold = {str(key).casefold(): value for key, value in configured.items()}
    replacements: dict[str, Any] = {}
    missing: list[str] = []
    for marker in markers:
        name = str(marker["name"])
        mapping = configured.get(name, configured_casefold.get(name.casefold(), name))
        required = True
        found = False
        value: Any = None
        if isinstance(mapping, dict):
            required = bool(mapping.get("required", True))
            if "value" in mapping:
                found, value = True, mapping["value"]
            else:
                sources = [str(candidate) for candidate in (mapping.get("field"), mapping.get("source")) if candidate]
                if not sources:
                    sources = [name]
                for source in sources:
                    found, value = _lookup(context, source)
                    if found and value not in (None, "", []):
                        break
        elif isinstance(mapping, str):
            if mapping.startswith(("literal:", "value:")):
                found, value = True, mapping.split(":", 1)[1]
            else:
                found, value = _lookup(context, mapping)
                if not found and mapping != name and "." not in mapping:
                    found, value = True, mapping
        else:
            found, value = True, mapping
        if not found or value in (None, "", []):
            if required:
                missing.append(name)
            else:
                replacements[name] = ""
        else:
            replacements[name] = value
    return replacements, sorted(set(missing), key=str.casefold)
