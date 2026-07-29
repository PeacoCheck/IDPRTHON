from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document as WordDocument

from backend.docx_templates import (
    MARKER_PATTERN,
    inspect_docx_markers,
    iter_document_paragraphs,
    render_docx_template,
    resolve_marker_values,
)


class DocxTemplateTests(unittest.TestCase):
    def test_markers_are_replaced_in_runs_tables_headers_and_footers(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source = root / "source.docx"
            target = root / "target.docx"
            document = WordDocument()
            paragraph = document.add_paragraph("Заказчик: ")
            first_marker_run = paragraph.add_run("[[za")
            first_marker_run.bold = True
            second_marker_run = paragraph.add_run("kaz]]")
            second_marker_run.bold = True
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Адрес: {{address}}"
            document.sections[0].header.paragraphs[0].text = "Журнал <<number>>"
            document.sections[0].footer.paragraphs[0].text = "Опция [[optional]]"
            document.save(source)

            markers = inspect_docx_markers(source)
            self.assertEqual({"address", "number", "optional", "zakaz"}, {item["name"] for item in markers})
            replacements, missing = resolve_marker_values(
                markers,
                {
                    "zakaz": {"field": "customer.name", "source": "object.meta.customer.name"},
                    "address": "object.address",
                    "number": "document.number",
                    "optional": {"source": "missing.optional", "required": False},
                },
                {
                    "customer.name": "ООО «Ручное значение»",
                    "object": {"address": "Самара", "meta": {"customer": {"name": "Автозначение"}}},
                    "document": {"number": "ОЖР-17"},
                },
            )
            self.assertEqual([], missing)
            replaced = render_docx_template(source, target, replacements)
            self.assertEqual(4, replaced)

            rendered = WordDocument(target)
            all_text = "\n".join(paragraph.text for _location, paragraph in iter_document_paragraphs(rendered))
            self.assertIn("ООО «Ручное значение»", all_text)
            self.assertIn("Самара", all_text)
            self.assertIn("ОЖР-17", all_text)
            self.assertIsNone(MARKER_PATTERN.search(all_text))
            replacement_runs = [run for run in rendered.paragraphs[0].runs if "Ручное значение" in run.text]
            self.assertEqual(1, len(replacement_runs))
            self.assertTrue(replacement_runs[0].bold)

    def test_required_unknown_marker_is_reported(self) -> None:
        markers = [{"name": "unknown", "tokens": ["[[unknown]]"], "occurrences": 1, "locations": []}]
        replacements, missing = resolve_marker_values(markers, {"unknown": {"field": "missing.value"}}, {})
        self.assertEqual({}, replacements)
        self.assertEqual(["unknown"], missing)


if __name__ == "__main__":
    unittest.main()
